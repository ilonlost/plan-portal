from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from app.models.entities import LineCapability, Product, ProductionLine


ARTICLE_HEADERS = [
    "ID (не менять)", "Артикул", "Продукция", "Тип", "Статус даты", "Статус арт",
    "Вес единицы, кг", "Штук в коробе", "Вес короба, кг", "Состояние записи",
]
PARAMETER_HEADERS = [
    "ID связи (не менять)", "Артикул", "Код линии (не менять)", "Цех", "Линия",
    "Скорость, кг/ч", "Квант замеса, кг", "Минимальный заказ, кг", "Короб",
    "Монопродукт", "Ограничения",
]


def _decimal(value, label: str, *, positive: bool = False) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        result = Decimal(str(value).replace(" ", "").replace(",", "."))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{label}: ожидается число") from exc
    if positive and result <= 0:
        raise ValueError(f"{label}: значение должно быть больше нуля")
    if not positive and result < 0:
        raise ValueError(f"{label}: значение не может быть отрицательным")
    return result


def _style_sheet(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    fill = PatternFill("solid", fgColor="D80C35")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.column_dimensions["A"].hidden = True


def export_catalog(db: Session) -> bytes:
    products = list(db.scalars(
        select(Product).where(Product.catalog_status != "deleted").order_by(Product.sku)
    ))
    capabilities = list(db.scalars(
        select(LineCapability).join(LineCapability.product).join(LineCapability.line)
        .where(Product.catalog_status != "deleted", ProductionLine.status == "active")
        .options(joinedload(LineCapability.product), joinedload(LineCapability.line))
        .order_by(Product.sku, ProductionLine.name)
    ))
    workbook = Workbook()
    info = workbook.active
    info.title = "Инструкция"
    info.append(["Справочник PLAN Portal"])
    info.append(["Редактируйте значения на листах «Артикулы» и «Параметры SKU», затем загрузите этот же файл обратно."])
    info.append(["Строки не удаляются из БД автоматически. Для блокировки или удаления измените «Состояние записи» на blocked или deleted."])
    info.append(["Допустимые состояния: active, blocked, deleted. Тип продукции: ОХЛ или ЗАМ."])
    info.column_dimensions["A"].width = 125
    info["A1"].font = Font(bold=True, size=16, color="D80C35")

    articles = workbook.create_sheet("Артикулы")
    articles.append(ARTICLE_HEADERS)
    for product in products:
        articles.append([
            product.id, product.sku, product.name, product.state, product.advance_status, product.fk_status,
            product.unit_weight_kg, product.units_per_box, product.box_weight_kg, product.catalog_status,
        ])
    _style_sheet(articles, [16, 18, 54, 14, 18, 22, 18, 18, 18, 20])
    state_validation = DataValidation(type="list", formula1='"active,blocked,deleted"')
    type_validation = DataValidation(type="list", formula1='"ОХЛ,ЗАМ"')
    articles.add_data_validation(state_validation)
    articles.add_data_validation(type_validation)
    state_validation.add(f"J2:J{max(2, articles.max_row + 100)}")
    type_validation.add(f"D2:D{max(2, articles.max_row + 100)}")

    params = workbook.create_sheet("Параметры SKU")
    params.append(PARAMETER_HEADERS)
    for item in capabilities:
        params.append([
            item.id, item.product.sku, item.line.code, item.line.workshop_name, item.line.name,
            item.units_per_hour, item.batch_quantum_kg, item.min_order_kg, item.product.units_per_box,
            item.product.mono_group, item.restrictions,
        ])
    _style_sheet(params, [20, 18, 24, 16, 28, 18, 20, 23, 14, 24, 54])
    params.column_dimensions["C"].hidden = True
    target = BytesIO()
    workbook.save(target)
    return target.getvalue()


def _headers(sheet, expected: list[str]) -> dict[str, int]:
    values = [str(cell.value or "").strip() for cell in sheet[1]]
    missing = [name for name in expected if name not in values]
    if missing:
        raise ValueError(f"Лист «{sheet.title}»: отсутствуют столбцы: {', '.join(missing)}")
    return {name: values.index(name) for name in expected}


def _row_value(row, indexes: dict[str, int], name: str):
    return row[indexes[name]].value


def import_catalog(db: Session, content: bytes) -> dict:
    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
    except Exception as exc:
        raise ValueError("Не удалось открыть XLSX-файл справочника") from exc
    for required in ("Артикулы", "Параметры SKU"):
        if required not in workbook.sheetnames:
            raise ValueError(f"В файле отсутствует лист «{required}»")

    articles = workbook["Артикулы"]
    a = _headers(articles, ARTICLE_HEADERS)
    created_products = updated_products = created_capabilities = updated_capabilities = 0
    products_by_sku = {item.sku: item for item in db.scalars(select(Product))}
    products_by_id = {item.id: item for item in products_by_sku.values()}
    for excel_row, row in enumerate(articles.iter_rows(min_row=2), 2):
        sku = str(_row_value(row, a, "Артикул") or "").strip()
        if not sku:
            continue
        raw_id = _row_value(row, a, "ID (не менять)")
        product = products_by_id.get(int(raw_id)) if raw_id not in (None, "") else products_by_sku.get(sku)
        if product is None:
            product = Product(sku=sku, name=str(_row_value(row, a, "Продукция") or sku).strip(), active=True, catalog_status="active")
            db.add(product)
            db.flush()
            products_by_id[product.id] = product
            products_by_sku[sku] = product
            created_products += 1
        elif product.sku != sku and sku in products_by_sku:
            raise ValueError(f"Артикулы, строка {excel_row}: артикул {sku} уже существует")
        else:
            products_by_sku.pop(product.sku, None)
            products_by_sku[sku] = product
            updated_products += 1
        status = str(_row_value(row, a, "Состояние записи") or "active").strip().lower()
        if status not in {"active", "blocked", "deleted"}:
            raise ValueError(f"Артикулы, строка {excel_row}: состояние должно быть active, blocked или deleted")
        product.sku = sku
        product.name = str(_row_value(row, a, "Продукция") or sku).strip()
        product.state = str(_row_value(row, a, "Тип") or "").strip() or None
        product.advance_status = str(_row_value(row, a, "Статус даты") or "").strip() or None
        product.fk_status = str(_row_value(row, a, "Статус арт") or "").strip() or None
        product.unit_weight_kg = _decimal(_row_value(row, a, "Вес единицы, кг"), f"Артикулы, строка {excel_row}", positive=True)
        product.units_per_box = _decimal(_row_value(row, a, "Штук в коробе"), f"Артикулы, строка {excel_row}", positive=True)
        product.box_weight_kg = _decimal(_row_value(row, a, "Вес короба, кг"), f"Артикулы, строка {excel_row}", positive=True)
        product.catalog_status = status
        product.active = status == "active"

    params = workbook["Параметры SKU"]
    p = _headers(params, PARAMETER_HEADERS)
    for excel_row, row in enumerate(params.iter_rows(min_row=2), 2):
        sku = str(_row_value(row, p, "Артикул") or "").strip()
        if not sku:
            continue
        product = products_by_sku.get(sku)
        if not product:
            raise ValueError(f"Параметры SKU, строка {excel_row}: артикул {sku} отсутствует на листе «Артикулы»")
        raw_id = _row_value(row, p, "ID связи (не менять)")
        capability = db.get(LineCapability, int(raw_id)) if raw_id not in (None, "") else None
        line_code = str(_row_value(row, p, "Код линии (не менять)") or "").strip()
        line_name = str(_row_value(row, p, "Линия") or "").strip()
        line = db.scalar(select(ProductionLine).where(ProductionLine.code == line_code)) if line_code else None
        if not line and line_name:
            line = db.scalar(select(ProductionLine).where(ProductionLine.name == line_name, ProductionLine.status == "active"))
        if not line:
            raise ValueError(f"Параметры SKU, строка {excel_row}: производственная линия не найдена")
        if capability is None:
            capability = db.scalar(select(LineCapability).where(LineCapability.product_id == product.id, LineCapability.line_id == line.id))
        speed = _decimal(_row_value(row, p, "Скорость, кг/ч"), f"Параметры SKU, строка {excel_row}", positive=True)
        if speed is None:
            raise ValueError(f"Параметры SKU, строка {excel_row}: укажите скорость")
        if capability is None:
            capability = LineCapability(product_id=product.id, line_id=line.id, units_per_hour=speed)
            db.add(capability)
            created_capabilities += 1
        else:
            if capability.product_id != product.id:
                raise ValueError(f"Параметры SKU, строка {excel_row}: ID связи относится к другому артикулу")
            duplicate = db.scalar(select(LineCapability).where(
                LineCapability.product_id == product.id,
                LineCapability.line_id == line.id,
                LineCapability.id != capability.id,
            ))
            if duplicate:
                raise ValueError(f"Параметры SKU, строка {excel_row}: связь артикула {sku} и линии {line.name} уже существует")
            capability.line_id = line.id
            updated_capabilities += 1
        capability.units_per_hour = speed
        capability.batch_quantum_kg = _decimal(_row_value(row, p, "Квант замеса, кг"), f"Параметры SKU, строка {excel_row}", positive=True)
        capability.min_order_kg = _decimal(_row_value(row, p, "Минимальный заказ, кг"), f"Параметры SKU, строка {excel_row}")
        product.units_per_box = _decimal(_row_value(row, p, "Короб"), f"Параметры SKU, строка {excel_row}", positive=True)
        product.mono_group = str(_row_value(row, p, "Монопродукт") or "").strip() or None
        capability.restrictions = str(_row_value(row, p, "Ограничения") or "").strip() or None
    db.flush()
    return {
        "products_created": created_products, "products_updated": updated_products,
        "capabilities_created": created_capabilities, "capabilities_updated": updated_capabilities,
    }
