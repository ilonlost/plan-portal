"""Application XLSX export using the portal's existing openpyxl dependency."""
from collections import Counter
from decimal import Decimal
from datetime import date, datetime, timedelta
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.tail_buffer_service import CENTERS, FIRST_MOVEMENT_SCOPE, MOSCOW, iter_export_rows, iter_cycle_rows

MAX_DATA_ROWS = 1_048_575
HEADERS = ["Дата и время проводки (МСК)", "Цех", "МЗ", "Линия", "Артикул", "Продукция",
           "Дата создания SSCC", "Буфер назначения", "Название буфера", "SSCC",
           "Первое движение в выбранных месяцах (МСК)", "Буфер карточки SSCC", "Номер проводки", "Месяц источника", "Проверка", "Вес, кг (L52_MENGE_LE)", "Срок годности, суток"]
WIDTHS = [25, 9, 12, 24, 18, 56, 22, 19, 35, 26, 30, 22, 22, 19, 60, 22, 22]
CELL_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
HEADER_FILL = PatternFill("solid", fgColor="C8102E")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def append_row(sheet, values, header=False):
    cells = []
    for value in values:
        if isinstance(value, datetime) and value.tzinfo:
            value = value.astimezone(MOSCOW).replace(tzinfo=None)
        cell = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
            cell.data_type = "s"  # Text identifiers and untrusted source text must never become formulas.
            cell.number_format = "@"
        elif isinstance(value, datetime):
            cell.number_format = "dd.mm.yyyy hh:mm:ss"
        elif isinstance(value, date):
            cell.number_format = "dd.mm.yyyy"
        elif isinstance(value, (float, Decimal)):
            cell.number_format = "#,##0.00"
        elif isinstance(value, int):
            cell.number_format = "#,##0"
        cell.alignment = CELL_ALIGNMENT
        if header:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        cells.append(cell)
    sheet.append(cells)


def export_workbook(start: date, end: date, center: str = "", search: str = "") -> str:
    workbook = Workbook(write_only=True)
    summary = workbook.create_sheet("Сводка")
    for letter, width in [("A", 40), ("B", 85), ("C", 24), ("D", 18), ("E", 22)]:
        summary.column_dimensions[letter].width = width
    summary.sheet_format.defaultRowHeight = 30
    selected_centers = sum(not center or code == center for _, code, _ in CENTERS)
    for index in (17 + selected_centers, 18 + selected_centers):
        summary.row_dimensions[index].height = 64
    summary.freeze_panes = "A2"
    summary.row_dimensions[1].height = 28
    for index in (9, 10, 11):
        summary.row_dimensions[index].height = 48
    file = NamedTemporaryFile(prefix="plan-fact-", suffix=".xlsx", delete=False)
    path = file.name
    file.close()
    counts = Counter()
    weights = Counter()
    missing_weights = 0
    hours = {}
    warnings = 0
    sheet = None
    sheet_count = row_count = total = 0
    try:
        for row in iter_export_rows(start, end, center, search):
            if sheet is None or row_count >= MAX_DATA_ROWS:
                if sheet is not None:
                    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{row_count + 1}"
                sheet_count += 1
                sheet = workbook.create_sheet("Проводки" if sheet_count == 1 else f"Проводки {sheet_count}")
                sheet.freeze_panes = "A2"
                sheet.row_dimensions[1].height = 58
                sheet.sheet_format.defaultRowHeight = 32
                for index, width in enumerate(WIDTHS, 1):
                    sheet.column_dimensions[get_column_letter(index)].width = width
                append_row(sheet, HEADERS, header=True)
                row_count = 0
            append_row(sheet, [row["moved_at"], "ПЦ" if row["workshop_code"] == "PC" else "КЦ",
                row["source_center"], row["line_name"], row["sku"], row["product_name"], row["created_date"],
                row["target_buffer"], row["buffer_name"], row["sscc"], row["first_moved_at"],
                str(row["card_buffer"]) if row["card_buffer"] is not None else None,
                str(row["record_id"]), row["source_month"], row["warning"], row.get("weight_kg"), row.get("shelf_life_days")])
            counts[row["source_center"]] += 1
            weight = row.get("weight_kg")
            if weight is not None:
                weights[row["source_center"]] += Decimal(str(weight))
            else:
                missing_weights += 1
            hour = row["moved_at"].astimezone(MOSCOW).replace(minute=0, second=0, microsecond=0)
            bucket = hours.setdefault((row["source_center"], hour), dict(postings=0, weight_kg=None, missing=0))
            bucket["postings"] += 1
            if weight is None:
                bucket["missing"] += 1
            else:
                bucket["weight_kg"] = (bucket["weight_kg"] or Decimal(0)) + Decimal(str(weight))
            warnings += bool(row["warning"])
            row_count += 1
            total += 1
        if sheet is None:
            sheet = workbook.create_sheet("Проводки")
            for index, width in enumerate(WIDTHS, 1):
                sheet.column_dimensions[get_column_letter(index)].width = width
            sheet.freeze_panes = "A2"
            sheet.row_dimensions[1].height = 58
            append_row(sheet, HEADERS, header=True)
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{row_count + 1}"
        append_row(summary, ["Факт хвостовых буферов", "CSB DWH"], header=True)
        for values in (["Начало периода", start], ["Конец периода включительно", end],
                       ["Место затрат", center or "Все МЗ"], ["Поиск", search or "Не задан"],
                       ["Сформировано (МСК)", datetime.now(MOSCOW)], ["Проводок выгружено", total],
                       ["Строк с предупреждениями", warnings], ["Первое движение", FIRST_MOVEMENT_SCOPE],
                       ["Единицы", "Вес: L52_MENGE_LE, кг. Цикл разделяется разрывом проводок от 5 минут; это не подтверждённый простой оборудования."],
                       ["Источник", "CSB_FK_REP: cp_DWH_LA0052_YYYYMM, cp_DWH_SY8581, cp_DWH_SY0012_SY8212_SY9014_SY9118, cp_DWH_SY0315"], []):
            append_row(summary, values)
        append_row(summary, ["Цех", "Линия", "МЗ", "Проводок", "Вес, кг"], header=True)
        for workshop, code, name in CENTERS:
            if not center or center == code:
                append_row(summary, ["ПЦ" if workshop == "PC" else "КЦ", name, code, counts[code], weights.get(code, 0 if not counts[code] else None)])
        append_row(summary, [])
        append_row(summary, ["Вес выгруженных проводок, кг", sum(weights.values()) if total > missing_weights else (0 if not total else None)])
        append_row(summary, ["Проводок без веса", missing_weights])
        append_row(summary, ["Циклы", "Все артикулы выбранных МЗ, внутри периода. Поиск по артикулу/SSCC не меняет границы циклов."])
        append_row(summary, ["кг/ч", "Почасовой выпуск: сумма веса за календарный час (МСК). В циклах: вес / интервал от первой до последней проводки; при нулевом интервале не рассчитывается."])
        hourly = workbook.create_sheet("Почасовой выпуск")
        hourly.freeze_panes = "A2"
        hourly.row_dimensions[1].height = 48
        for index, width in enumerate([12, 24, 25, 25, 22, 20, 20, 22], 1):
            hourly.column_dimensions[get_column_letter(index)].width = width
        append_row(hourly, ["МЗ", "Линия", "Начало часа, МСК", "Конец часа, МСК", "Вес, кг", "Выпуск, кг/ч", "Проводок", "Проводок без веса"], header=True)
        for _, code, _ in CENTERS:
            if center and center != code:
                continue
            hour = datetime.combine(start, datetime.min.time(), tzinfo=MOSCOW)
            stop = datetime.combine(end + timedelta(days=1), datetime.min.time(), tzinfo=MOSCOW)
            while hour < stop:
                hours.setdefault((code, hour), dict(postings=0, weight_kg=Decimal(0), missing=0))
                hour += timedelta(hours=1)
        for (code, hour), bucket in sorted(hours.items()):
            name = next(name for _, c, name in CENTERS if c == code)
            append_row(hourly, [code, name, hour, hour + timedelta(hours=1), bucket["weight_kg"], bucket["weight_kg"], bucket["postings"], bucket["missing"]])
        hourly.auto_filter.ref = f"A1:H{len(hours) + 1}"
        cycles = workbook.create_sheet("Циклы")
        cycles.freeze_panes = "A2"
        cycles.row_dimensions[1].height = 58
        for index, width in enumerate([12, 24, 25, 25, 23, 22, 24, 24, 18, 23, 16], 1):
            cycles.column_dimensions[get_column_letter(index)].width = width
        append_row(cycles, ["МЗ", "Линия", "Начало цикла, МСК", "Конец цикла, МСК", "Интервал проводок, мин", "Вес, кг", "Вес / интервал, кг/ч", "Разрыв перед циклом, мин (≥5)", "Проводок", "Проводок без веса", "Цикл № на МЗ"], header=True)
        cycle_numbers = Counter()
        cycle_count = 0
        for cycle in iter_cycle_rows(start, end, center):
            cycle_numbers[cycle["source_center"]] += 1
            duration = Decimal(str((cycle["end_at"] - cycle["start_at"]).total_seconds()))
            rate = cycle["weight_kg"] * 3600 / duration if duration > 0 and cycle["weight_kg"] is not None else None
            name = next(name for _, c, name in CENTERS if c == cycle["source_center"])
            append_row(cycles, [cycle["source_center"], name, cycle["start_at"], cycle["end_at"], duration / 60,
                               cycle["weight_kg"], rate, cycle["gap_before_minutes"], cycle["postings"], cycle["missing_weight_count"], cycle_numbers[cycle["source_center"]]])
            cycle_count += 1
        cycles.auto_filter.ref = f"A1:K{cycle_count + 1}"
        workbook.save(path)
        return path
    except Exception:
        for sheet in workbook.worksheets:
            if not sheet.closed:
                sheet.close()
        # Saving closes/removes openpyxl's temporary worksheet XML files as well.
        try:
            workbook.save(path)
        except Exception:
            pass
        workbook.close()
        Path(path).unlink(missing_ok=True)
        raise
