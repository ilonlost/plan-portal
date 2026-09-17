from __future__ import annotations

from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session, joinedload

from app.core.security import UserContext, current_user, require_planner
from app.db.session import get_db
from app.models.entities import (
    AdvanceConfirmationBatch, AdvanceConfirmationItem, AuditEvent, DemandItem, ImportedOrder,
    LineCapability, Product, ProductionLine, ProductionScheduleItem,
)
from app.services.plan_service import PlanService


router = APIRouter(prefix="/advance-confirmations", tags=["advance-confirmations"])
MAX_FILE_SIZE = 30 * 1024 * 1024


class ConfirmationRow(BaseModel):
    row_number: int = Field(ge=1)
    sku: str = Field(min_length=1, max_length=80)
    product_name: str = Field(min_length=1, max_length=200)
    product_id: int | None = None
    line_id: int | None = None
    line_name: str = ""
    production_date: date
    marking_date: date
    current_quantity_kg: Decimal = Field(default=0, ge=0)
    quantity_kg: Decimal = Field(ge=0)
    actual_quantity_kg: Decimal = Field(default=0, ge=0)
    delta_quantity_kg: Decimal = 0
    advance_status: str | None = None
    valid: bool = True
    errors: list[str] = []
    warnings: list[str] = []


class ApplyRequest(BaseModel):
    file_name: str = Field(min_length=1, max_length=240)
    rows: list[ConfirmationRow]


def _decimal(value, default: Decimal | None = None) -> Decimal | None:
    if value in (None, ""):
        return default
    try:
        return Decimal(str(value).replace(" ", "").replace(",", ".")).quantize(Decimal("0.001"))
    except (InvalidOperation, ValueError):
        return None


def _filename_date(file_name: str) -> date:
    match = re.search(r"(?<!\d)(\d{1,2})[._-](\d{1,2})[._-](\d{4})(?!\d)", file_name)
    if not match:
        raise HTTPException(422, "В имени файла не найдена дата ДМ в формате ДД.ММ.ГГГГ")
    try:
        return date(int(match.group(3)), int(match.group(2)), int(match.group(1)))
    except ValueError as exc:
        raise HTTPException(422, "В имени файла указана некорректная дата ДМ") from exc


def _header(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().casefold())


def _sku(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "").strip().removesuffix(".0")


def _plan_matches(db: Session, product_id: int, marking_date: date) -> list[ProductionScheduleItem]:
    plan = PlanService(db).active_plan()
    if not plan:
        return []
    rows = list(db.scalars(
        select(ProductionScheduleItem).where(
            ProductionScheduleItem.plan_id == plan.id,
            ProductionScheduleItem.product_id == product_id,
            ProductionScheduleItem.schedule_kind == "production",
            ProductionScheduleItem.excluded.is_(False),
        ).options(joinedload(ProductionScheduleItem.line), joinedload(ProductionScheduleItem.demand_item))
    ))
    return [row for row in rows if row.marking_date == marking_date or (row.demand_item and row.demand_item.source_plan_date == marking_date)]


def _resolve_row(db: Session, row_number: int, sku: str, quantity: Decimal | None, actual: Decimal | None, marking_date: date) -> dict:
    errors: list[str] = []
    warnings: list[str] = []
    product = db.scalar(select(Product).where(Product.sku == sku, Product.catalog_status != "deleted")) if sku else None
    if not product:
        errors.append("Артикул отсутствует в справочнике")
    matches = _plan_matches(db, product.id, marking_date) if product else []
    current = sum((Decimal(item.quantity_kg or item.quantity or 0) for item in matches), Decimal("0"))
    line = next((item.line for item in matches if item.line), None)
    production_date = min((item.production_date for item in matches if item.production_date), default=None)
    if product and not line:
        capability = db.scalar(
            select(LineCapability).join(LineCapability.line).where(
                LineCapability.product_id == product.id, ProductionLine.status == "active",
            ).options(joinedload(LineCapability.line)).order_by(ProductionLine.priority, ProductionLine.id)
        )
        line = capability.line if capability else None
    if product and not line:
        errors.append("Для артикула не настроена производственная линия")
    if product and not matches:
        warnings.append("В действующем плане за эту ДМ задание не найдено; будет создана корректировка исходного спроса")
    if quantity is None or quantity < 0:
        errors.append("Некорректный подтверждённый объём")
        quantity = Decimal("0")
    if actual is None or actual < 0:
        errors.append("Некорректный фактический объём")
        actual = Decimal("0")
    advance = (product.advance_status or "").strip().upper() == "АЗ" if product else False
    production_date = marking_date - timedelta(days=1) if advance else production_date or marking_date
    return {
        "row_number": row_number, "sku": sku, "product_name": product.name if product else "Не найдено",
        "product_id": product.id if product else None, "line_id": line.id if line else None,
        "line_name": line.name if line else "", "production_date": production_date, "marking_date": marking_date,
        "current_quantity_kg": current, "quantity_kg": quantity, "actual_quantity_kg": actual,
        "delta_quantity_kg": quantity - current, "advance_status": product.advance_status if product else None,
        "valid": not errors, "errors": errors, "warnings": warnings,
    }


@router.post("/preview")
async def preview_confirmation(
    file: UploadFile = File(...), db: Session = Depends(get_db), user: UserContext = Depends(require_planner),
) -> dict:
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(422, "Загрузите файл Excel в формате XLSX или XLSM")
    marking_date = _filename_date(file.filename)
    content = await file.read(MAX_FILE_SIZE + 1)
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(413, "Размер файла превышает 30 МБ")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(422, "Не удалось прочитать файл Excel") from exc
    sheet = workbook.active
    headers = {_header(cell.value): index for index, cell in enumerate(next(sheet.iter_rows(min_row=1, max_row=1)), start=1)}
    material_col = headers.get("материал")
    quantity_col = headers.get("всего кг")
    actual_col = headers.get("сколько произвели")
    if not material_col or not quantity_col:
        raise HTTPException(422, "В файле нужны столбцы «Материал» и «Всего КГ»")
    rows = []
    for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        sku = _sku(values[material_col - 1] if len(values) >= material_col else None)
        if not sku:
            continue
        quantity = _decimal(values[quantity_col - 1] if len(values) >= quantity_col else None)
        actual = _decimal(values[actual_col - 1] if actual_col and len(values) >= actual_col else None, quantity)
        rows.append(_resolve_row(db, number, sku, quantity, actual, marking_date))
    if not rows:
        raise HTTPException(422, "В файле не найдено ни одной позиции")
    current_total = sum((row["current_quantity_kg"] for row in rows), Decimal("0"))
    quantity_total = sum((row["quantity_kg"] for row in rows), Decimal("0"))
    return {
        "file_name": file.filename, "sheet_name": sheet.title, "marking_date": marking_date,
        "rows": rows, "summary": {
            "total_rows": len(rows), "valid_rows": sum(1 for row in rows if row["valid"]),
            "invalid_rows": sum(1 for row in rows if not row["valid"]),
            "current_quantity_kg": current_total, "quantity_kg": quantity_total,
            "delta_quantity_kg": quantity_total - current_total,
        },
    }


def _latest_ohl_order(db: Session) -> ImportedOrder | None:
    return db.scalar(select(ImportedOrder).where(ImportedOrder.template_type == "ohl_daily").order_by(ImportedOrder.imported_at.desc(), ImportedOrder.id.desc()))


@router.post("/apply")
def apply_confirmation(
    payload: ApplyRequest, db: Session = Depends(get_db), user: UserContext = Depends(require_planner),
) -> dict:
    if not payload.rows:
        raise HTTPException(422, "Нет строк для применения")
    order = _latest_ohl_order(db)
    plan_service = PlanService(db)
    plan = plan_service.active_plan()
    if not order or not plan:
        raise HTTPException(409, "Сначала загрузите основной план ОХЛ и сформируйте действующий план")
    batch_date = payload.rows[0].marking_date
    batch = AdvanceConfirmationBatch(
        file_name=payload.file_name, marking_date=batch_date, status="applying", total_rows=len(payload.rows),
        total_quantity_kg=sum((row.quantity_kg for row in payload.rows), Decimal("0")),
        total_actual_kg=sum((row.actual_quantity_kg for row in payload.rows), Decimal("0")), created_by=user.username,
    )
    db.add(batch)
    db.flush()
    max_source_row = db.scalar(select(func.max(DemandItem.source_row)).where(DemandItem.order_id == order.id)) or 1
    touched: list[tuple[AdvanceConfirmationItem, list[int]]] = []
    for index, row in enumerate(payload.rows, start=1):
        product = db.scalar(select(Product).where(Product.sku == row.sku, Product.catalog_status != "deleted"))
        if not product:
            raise HTTPException(422, f"Строка {row.row_number}: артикул {row.sku} отсутствует в справочнике")
        line = db.get(ProductionLine, row.line_id) if row.line_id else None
        capability = db.scalar(select(LineCapability).where(LineCapability.product_id == product.id, LineCapability.line_id == row.line_id)) if line else None
        if not line or line.status != "active" or not capability:
            raise HTTPException(422, f"Строка {row.row_number}: выбранная линия не связана с артикулом {row.sku}")
        matches = _plan_matches(db, product.id, row.marking_date)
        previous = sum((Decimal(item.quantity_kg or item.quantity or 0) for item in matches), Decimal("0"))
        demands = list(db.scalars(select(DemandItem).where(
            DemandItem.order_id == order.id, DemandItem.product_id == product.id,
            DemandItem.source_plan_date == row.marking_date,
        ).order_by(DemandItem.id)))
        if not demands:
            demand = DemandItem(
                order_id=order.id, product=product, source_row=max_source_row + index, sku=product.sku,
                product_name=product.name, quantity=row.quantity_kg, source_quantity=row.quantity_kg,
                source_unit="кг", quantity_kg=row.quantity_kg, source_kind="ohl", source_plan_date=row.marking_date,
                marking_date=row.marking_date, requested_date=row.production_date, due_date=row.production_date,
                exact_date=True, advance_production=(product.advance_status or "").strip().upper() == "АЗ",
                raw_data={}, valid=True, validation_errors=[],
            )
            db.add(demand)
            db.flush()
            demands = [demand]
        for demand_index, demand in enumerate(demands):
            corrected = row.quantity_kg if demand_index == 0 else Decimal("0")
            demand.quantity = corrected
            demand.quantity_kg = corrected
            demand.source_quantity = corrected
            demand.source_unit = "кг"
            demand.source_plan_date = row.marking_date
            demand.marking_date = row.marking_date
            demand.requested_date = row.production_date
            demand.due_date = row.production_date
            demand.exact_date = True
            demand.valid = True
            demand.raw_data = {
                **(demand.raw_data or {}), "preferred_line_id": line.id,
                "az_override_production_date": row.production_date.isoformat(), "advance_confirmation_batch_id": batch.id,
            }
        history = AdvanceConfirmationItem(
            batch=batch, source_row=row.row_number, product=product, line=line, sku=row.sku,
            product_name=row.product_name.strip() or product.name, line_name=line.name,
            production_date=row.production_date, marking_date=row.marking_date,
            previous_quantity_kg=previous, quantity_kg=row.quantity_kg, actual_quantity_kg=row.actual_quantity_kg,
            delta_quantity_kg=row.quantity_kg - previous, advance_status=product.advance_status,
            demand_item_ids=[item.id for item in demands], warnings=row.warnings,
        )
        db.add(history)
        touched.append((history, [item.id for item in demands]))
    db.flush()
    demands = plan_service._latest_source_demands()
    plan_service.calculate(plan, demands, "advance_confirmation_applied")
    for history, demand_ids in touched:
        schedule = list(db.scalars(select(ProductionScheduleItem).where(
            ProductionScheduleItem.plan_id == plan.id, ProductionScheduleItem.demand_item_id.in_(demand_ids),
            ProductionScheduleItem.schedule_kind == "production", ProductionScheduleItem.excluded.is_(False),
        )))
        history.schedule_item_ids = [item.id for item in schedule]
        planned_total = sum((Decimal(item.quantity_kg or item.quantity or 0) for item in schedule), Decimal("0"))
        remaining_actual = Decimal(history.actual_quantity_kg)
        for item_index, item in enumerate(schedule):
            if item_index == len(schedule) - 1:
                allocated = remaining_actual
            else:
                allocated = (Decimal(history.actual_quantity_kg) * Decimal(item.quantity_kg or item.quantity or 0) / planned_total).quantize(Decimal("0.001")) if planned_total else Decimal("0")
                remaining_actual -= allocated
            item.actual_quantity_kg = max(Decimal("0"), allocated)
    batch.status = "applied"
    db.add(AuditEvent(
        username=user.username, action="advance_confirmation_applied", entity_type="advance_confirmation_batch", entity_id=str(batch.id),
        details={"file_name": batch.file_name, "marking_date": batch.marking_date.isoformat(), "rows": batch.total_rows,
                 "quantity_kg": str(batch.total_quantity_kg), "actual_quantity_kg": str(batch.total_actual_kg), "plan_id": plan.id},
    ))
    db.commit()
    return {"ok": True, "batch_id": batch.id, "plan_id": plan.id, "plan_recalculated": True, "rows": batch.total_rows}


def _item_dict(item: AdvanceConfirmationItem) -> dict:
    return {
        "id": item.id, "row_number": item.source_row, "sku": item.sku, "product_name": item.product_name,
        "line_id": item.line_id, "line_name": item.line_name, "production_date": item.production_date,
        "marking_date": item.marking_date, "current_quantity_kg": item.previous_quantity_kg,
        "quantity_kg": item.quantity_kg, "actual_quantity_kg": item.actual_quantity_kg,
        "delta_quantity_kg": item.delta_quantity_kg, "advance_status": item.advance_status,
        "schedule_item_ids": item.schedule_item_ids, "warnings": item.warnings,
    }


@router.get("")
def confirmation_history(
    limit: int = 20, db: Session = Depends(get_db), user: UserContext = Depends(current_user),
) -> list[dict]:
    batches = list(db.scalars(
        select(AdvanceConfirmationBatch).options(joinedload(AdvanceConfirmationBatch.items))
        .order_by(AdvanceConfirmationBatch.created_at.desc(), AdvanceConfirmationBatch.id.desc()).limit(min(max(limit, 1), 100))
    ).unique())
    return [{
        "id": batch.id, "file_name": batch.file_name, "marking_date": batch.marking_date, "status": batch.status,
        "total_rows": batch.total_rows, "total_quantity_kg": batch.total_quantity_kg,
        "total_actual_kg": batch.total_actual_kg, "created_by": batch.created_by, "created_at": batch.created_at,
        "items": [_item_dict(item) for item in sorted(batch.items, key=lambda value: value.source_row)],
    } for batch in batches]


@router.get("/{batch_id}/export.xlsx")
def export_confirmation(
    batch_id: int, db: Session = Depends(get_db), user: UserContext = Depends(current_user),
) -> Response:
    batch = db.scalar(select(AdvanceConfirmationBatch).where(AdvanceConfirmationBatch.id == batch_id).options(joinedload(AdvanceConfirmationBatch.items)))
    if not batch:
        raise HTTPException(404, "Корректировка АЗ не найдена")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Факт выполнения АЗ"
    headers = ["Материал", "Продукция", "Линия", "ДП", "ДМ", "План до, кг", "Подтверждено, кг", "Изменение, кг", "Произведено, кг", "Статус АЗ"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="D71035")
        cell.alignment = Alignment(horizontal="center")
    plan = PlanService(db).active_plan()
    for item in sorted(batch.items, key=lambda value: value.source_row):
        current_actual = None
        if plan and item.demand_item_ids:
            actual_values = list(db.scalars(select(ProductionScheduleItem.actual_quantity_kg).where(
                ProductionScheduleItem.plan_id == plan.id,
                ProductionScheduleItem.demand_item_id.in_(item.demand_item_ids),
                ProductionScheduleItem.schedule_kind == "production",
                ProductionScheduleItem.excluded.is_(False),
                ProductionScheduleItem.actual_quantity_kg.is_not(None),
            )))
            if actual_values:
                current_actual = sum((Decimal(value) for value in actual_values), Decimal("0"))
        sheet.append([item.sku, item.product_name, item.line_name, item.production_date, item.marking_date,
                      float(item.previous_quantity_kg), float(item.quantity_kg), float(item.delta_quantity_kg),
                      float(current_actual if current_actual is not None else item.actual_quantity_kg), item.advance_status or ""])
    for column, width in {"A": 16, "B": 55, "C": 25, "D": 13, "E": 13, "F": 16, "G": 18, "H": 16, "I": 18, "J": 14}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    filename = f"AZ_fact_{batch.marking_date.strftime('%d.%m.%Y')}.xlsx"
    return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})
