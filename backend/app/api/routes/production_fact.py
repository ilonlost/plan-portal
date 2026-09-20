from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, time, timedelta
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.core.security import UserContext, current_user


router = APIRouter(prefix="/production-fact", tags=["production-fact"])
MAX_RANGE_DAYS = 184

PRODUCTS = {
    "PC": [
        ("1010035972", "Булочка пшеничная"),
        ("1010027200", "Круассан сливочный"),
        ("1010027261", "Булочка солодовая"),
        ("1010026337", "Багет французский"),
    ],
    "KC": [
        ("1010035951", "Чиабатта с ветчиной"),
        ("1010022357", "Сырники классические"),
        ("1010034117", "Лазанья Болоньезе"),
        ("1010035946", "Чизбургер"),
        ("1010028780", "Борщ со сметаной"),
    ],
}

CENTERS = [
    {"code": "5810", "workshop": "PC", "name": "Булка", "line": "ПЦ · Булка", "rate": 8.4},
    {"code": "5800", "workshop": "PC", "name": "Слойка", "line": "ПЦ · Слойка", "rate": 7.2},
    {"code": "5820", "workshop": "PC", "name": "Хлеб", "line": "ПЦ · Хлеб", "rate": 9.1},
    {"code": "5830", "workshop": "PC", "name": "Ручная зона", "line": "ПЦ · Ручная зона", "rate": 4.6},
    {"code": "5480", "workshop": "KC", "name": "Сэндвичи", "line": "КЦ · Сэндвичи", "rate": 6.1},
    {"code": "5410", "workshop": "KC", "name": "Жареные блюда", "line": "КЦ · Жареные блюда", "rate": 7.7},
    {"code": "5440", "workshop": "KC", "name": "Лазанья", "line": "КЦ · Лазанья", "rate": 5.8},
    {"code": "5460", "workshop": "KC", "name": "Бургеры", "line": "КЦ · Бургеры", "rate": 6.8},
    {"code": "5430", "workshop": "KC", "name": "Супы", "line": "КЦ · Супы", "rate": 8.0},
]

KIND_LABELS = {
    "production": "Производство",
    "pause": "Пауза",
    "changeover": "Переналадка",
    "maintenance": "ТО",
}


def _range(start: date | None, end: date | None) -> tuple[date, date, int]:
    start = start or date.today()
    end = end or start
    if end < start:
        raise HTTPException(422, "Конечная дата не может быть раньше начальной")
    days = (end - start).days + 1
    if days > MAX_RANGE_DAYS:
        raise HTTPException(422, "Можно выбрать период не более шести месяцев (184 дня)")
    return start, end, days


def _iso(day: date, minute: int) -> str:
    return datetime.combine(day, time.min).replace(hour=minute // 60, minute=minute % 60).isoformat()


def _event(center: dict, day: date, index: int, sequence: int, kind: str, start_minute: int, duration: int, product_index: int) -> dict:
    products = PRODUCTS[center["workshop"]]
    sku, product_name = products[product_index % len(products)]
    production = kind == "production"
    quantity = round(duration * center["rate"] * (0.96 + ((day.toordinal() + index) % 7) / 100), 2) if production else None
    status = "completed" if day < date.today() else ("in_progress" if production and sequence == 3 else "completed")
    return {
        "id": f"stub-{center['code']}-{day.isoformat()}-{sequence}",
        "date": day.isoformat(),
        "workshop_code": center["workshop"],
        "cost_center_code": center["code"],
        "cost_center_name": center["name"],
        "line_name": center["line"],
        "kind": kind,
        "kind_label": KIND_LABELS[kind],
        "start_at": _iso(day, start_minute),
        "end_at": _iso(day, start_minute + duration),
        "duration_minutes": duration,
        "sku": sku if production else None,
        "product_name": product_name if production else ("Регламентная переналадка" if kind == "changeover" else "Остановка оборудования"),
        "quantity_kg": quantity,
        "status": status,
        "source": "erp_stub",
    }


def _events_for_center(center: dict, start: date, end: date) -> list[dict]:
    index = CENTERS.index(center)
    events: list[dict] = []
    day = start
    while day <= end:
        seed = day.toordinal() + index * 17
        first_start = 350 + seed % 41
        first_duration = 205 + seed % 66
        pause_duration = 18 + seed % 24
        second_duration = 185 + (seed * 3) % 86
        product_index = seed + index
        cursor = first_start
        events.append(_event(center, day, index, 1, "production", cursor, first_duration, product_index))
        cursor += first_duration
        events.append(_event(center, day, index, 2, "pause", cursor, pause_duration, product_index))
        cursor += pause_duration
        events.append(_event(center, day, index, 3, "production", cursor, second_duration, product_index + 1))
        cursor += second_duration
        if seed % 3 != 0:
            events.append(_event(center, day, index, 4, "changeover", cursor, 30 + seed % 16, product_index + 1))
            cursor += 30 + seed % 16
            events.append(_event(center, day, index, 5, "production", cursor, 95 + seed % 51, product_index + 2))
        elif seed % 11 == 0:
            events.append(_event(center, day, index, 4, "maintenance", cursor, 60, product_index + 1))
        day += timedelta(days=1)
    return events


def _center_summary(center: dict, start: date, end: date, include_timeline: bool) -> dict:
    events = _events_for_center(center, start, end)
    grouped: dict[str, dict] = {}
    for event in events:
        item = grouped.setdefault(event["date"], {"date": event["date"], "worked_hours": 0.0, "pause_hours": 0.0, "quantity_kg": 0.0})
        if event["kind"] == "production":
            item["worked_hours"] += event["duration_minutes"] / 60
            item["quantity_kg"] += event["quantity_kg"] or 0
        else:
            item["pause_hours"] += event["duration_minutes"] / 60
    daily = []
    for value in grouped.values():
        daily.append({**value, "worked_hours": round(value["worked_hours"], 2), "pause_hours": round(value["pause_hours"], 2), "quantity_kg": round(value["quantity_kg"], 2)})
    worked = sum(value["worked_hours"] for value in daily)
    paused = sum(value["pause_hours"] for value in daily)
    return {
        "code": center["code"], "name": center["name"], "workshop_code": center["workshop"], "line_name": center["line"],
        "worked_hours": round(worked, 2), "pause_hours": round(paused, 2),
        "utilization_percent": round(min(100, worked / max(1, len(daily) * 16) * 100), 1),
        "daily": daily, "timeline_events": events if include_timeline else [],
    }


def _articles(all_events: list[dict]) -> list[dict]:
    rows: dict[tuple[str, str, str], dict] = {}
    for event in all_events:
        if event["kind"] != "production":
            continue
        key = (event["sku"], event["product_name"], event["workshop_code"])
        row = rows.setdefault(key, {"sku": key[0], "product_name": key[1], "workshop_code": key[2], "total_hours": 0.0, "total_quantity_kg": 0.0, "days": defaultdict(lambda: {"hours": 0.0, "quantity_kg": 0.0})})
        row["total_hours"] += event["duration_minutes"] / 60
        row["total_quantity_kg"] += event["quantity_kg"] or 0
        day = row["days"][event["date"]]
        day["hours"] += event["duration_minutes"] / 60
        day["quantity_kg"] += event["quantity_kg"] or 0
    result = []
    for row in rows.values():
        result.append({
            **{key: value for key, value in row.items() if key != "days"},
            "total_hours": round(row["total_hours"], 2), "total_quantity_kg": round(row["total_quantity_kg"], 2),
            "days": [{"date": day, "hours": round(values["hours"], 2), "quantity_kg": round(values["quantity_kg"], 2)} for day, values in sorted(row["days"].items())],
        })
    return sorted(result, key=lambda item: item["total_hours"], reverse=True)


def _process_maps(end: date) -> list[dict]:
    result = []
    for index, center in enumerate(CENTERS):
        product = PRODUCTS[center["workshop"]][(end.toordinal() + index) % len(PRODUCTS[center["workshop"]])]
        if center["workshop"] == "PC":
            stage_names = ["МЗ создано", "Начало перетарки", f"Цех {center['name'].lower()}", f"Маркировка {center['name'].lower()}", "Передача на склад"]
        else:
            stage_names = ["МЗ создано", "Подготовка сырья", f"Производство · {center['name']}", "Охлаждение", "Маркировка и склад"]
        progress = 38 + ((end.toordinal() + index * 13) % 60)
        stages = []
        for stage_index, name in enumerate(stage_names):
            threshold = stage_index * 25
            status = "completed" if progress >= threshold + 25 else ("active" if progress >= threshold else "pending")
            stages.append({"id": f"{center['code']}-{stage_index}", "name": name, "status": status, "progress_percent": 100 if status == "completed" else (max(8, (progress - threshold) * 4) if status == "active" else 0)})
        result.append({"workshop_code": center["workshop"], "line_name": center["line"], "cost_center_code": center["code"], "sku": product[0], "product_name": product[1], "progress_percent": progress, "current_stage": next((stage["name"] for stage in stages if stage["status"] == "active"), "Выполнено"), "stages": stages})
    return result


def _payload(start: date, end: date, days: int) -> dict:
    all_events: list[dict] = []
    centers = []
    for center in CENTERS:
        events = _events_for_center(center, start, end)
        all_events.extend(events)
        centers.append(_center_summary(center, start, end, days <= 7))
    workshops = []
    for code, name in (("PC", "ПЦ · Пекарный цех"), ("KC", "КЦ · Кулинарный цех")):
        workshop_centers = [center for center in centers if center["workshop_code"] == code]
        workshops.append({"code": code, "name": name, "total_hours": round(sum(center["worked_hours"] for center in workshop_centers), 2), "pause_hours": round(sum(center["pause_hours"] for center in workshop_centers), 2), "cost_centers": workshop_centers})
    production_events = [event for event in all_events if event["kind"] == "production"]
    return {
        "source": "erp_stub", "generated_at": datetime.now().isoformat(),
        "range": {"start": start.isoformat(), "end": end.isoformat(), "days": days, "max_days": MAX_RANGE_DAYS},
        "summary": {"production_hours": round(sum(event["duration_minutes"] for event in production_events) / 60, 2), "pause_hours": round(sum(event["duration_minutes"] for event in all_events if event["kind"] != "production") / 60, 2), "quantity_kg": round(sum(event["quantity_kg"] or 0 for event in production_events), 2), "cost_centers": len(CENTERS), "operations": len(all_events)},
        "workshops": workshops, "articles": _articles(all_events), "process_maps": _process_maps(end),
    }


@router.get("")
def production_fact(start: date | None = None, end: date | None = None, _: UserContext = Depends(current_user)) -> dict:
    range_start, range_end, days = _range(start, end)
    return _payload(range_start, range_end, days)


@router.get("/cost-centers/{code}")
def production_fact_detail(code: str, start: date | None = None, end: date | None = None, _: UserContext = Depends(current_user)) -> dict:
    range_start, range_end, days = _range(start, end)
    center = next((item for item in CENTERS if item["code"] == code), None)
    if not center:
        raise HTTPException(404, "Место затрат не найдено")
    events = _events_for_center(center, range_start, range_end)
    return {"source": "erp_stub", "range": {"start": range_start.isoformat(), "end": range_end.isoformat(), "days": days}, "cost_center": _center_summary(center, range_start, range_end, False), "events": events}


@router.get("/export.xlsx")
def export_production_fact(start: date | None = None, end: date | None = None, _: UserContext = Depends(current_user)) -> Response:
    range_start, range_end, days = _range(start, end)
    payload = _payload(range_start, range_end, days)
    workbook = Workbook()
    events_sheet = workbook.active
    events_sheet.title = "События"
    headers = ["Дата", "Цех", "МЗ", "Линия", "Тип", "Начало", "Окончание", "Длительность, ч", "Артикул", "Продукция", "Количество, кг", "Источник"]
    events_sheet.append(headers)
    for center in CENTERS:
        for event in _events_for_center(center, range_start, range_end):
            events_sheet.append([event["date"], "ПЦ" if event["workshop_code"] == "PC" else "КЦ", event["cost_center_code"], event["line_name"], event["kind_label"], event["start_at"][11:16], event["end_at"][11:16], round(event["duration_minutes"] / 60, 2), event["sku"] or "", event["product_name"] or "", event["quantity_kg"] or "", "Заглушка ERP/CSB"])
    article_sheet = workbook.create_sheet("Артикулы ГП")
    article_sheet.append(["Артикул", "Продукция", "Цех", "Время, ч", "Количество, кг"])
    for article in payload["articles"]:
        article_sheet.append([article["sku"], article["product_name"], "ПЦ" if article["workshop_code"] == "PC" else "КЦ", article["total_hours"], article["total_quantity_kg"]])
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="D80C35")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(42, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    output = BytesIO()
    workbook.save(output)
    filename = f"Факт производства {range_start:%d.%m.%Y}-{range_end:%d.%m.%Y}.xlsx"
    return Response(content=output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})
