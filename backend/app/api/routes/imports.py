from datetime import timedelta
from decimal import Decimal

from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.session import get_db
from app.core.security import UserContext, require_planner
from app.core.line_catalog import line_key, workshop_for_line
from app.models.entities import (
    AuditEvent, DemandItem, ImportedOrder, ImportFile, LineCapability, LineCapacity, Product,
    ProductionLine, ProductionPlan, ProductionScheduleItem, ScheduleStatus,
)
from app.schemas.common import ImportConfirmRequest, ImportPreview
from app.services.import_service import ExcelImportService
from app.services.line_schedule_service import ensure_line_capacities
from app.services.plan_service import PlanService, plan_dict
from app.services.notification_service import send_notification

router = APIRouter(prefix="/imports", tags=["imports"])


@router.get("/maintenance-template.xlsx")
def maintenance_template(db: Session = Depends(get_db), user: UserContext = Depends(require_planner)) -> StreamingResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "График ТО"
    headers = ["Линия", "Дата", "Смена", "Тип события", "Длительность, ч", "Причина / работы"]
    sheet.append(headers)
    red = "D90B32"
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor=red)
        cell.alignment = Alignment(horizontal="center")
    lines = list(db.scalars(select(ProductionLine).where(ProductionLine.status == "active").order_by(ProductionLine.workshop_code, ProductionLine.priority, ProductionLine.name)))
    for line in lines:
        sheet.append([line.name, None, "День", "ТО", None, "Плановое техническое обслуживание"])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:F{max(2, len(lines) + 1)}"
    widths = [28, 14, 13, 17, 18, 48]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in range(2, max(1002, len(lines) + 2)):
        sheet.cell(row, 2).number_format = "DD.MM.YYYY"
    shift_validation = DataValidation(type="list", formula1='"День,Ночь"')
    type_validation = DataValidation(type="list", formula1='"ТО,Мойка,Простой"')
    sheet.add_data_validation(shift_validation); shift_validation.add("C2:C1001")
    sheet.add_data_validation(type_validation); type_validation.add("D2:D1001")
    guide = workbook.create_sheet("Инструкция")
    guide.append(["Шаблон графика ТО"])
    guide.append(["Заполните дату, смену, тип события, длительность и описание работ. Название линии не меняйте."])
    guide.append(["Повторная загрузка заменяет только события из предыдущего файла ТО. Текущий производственный план и ручные события сохраняются."])
    guide.column_dimensions["A"].width = 115
    guide["A1"].font = Font(bold=True, color=red, size=14)
    output = BytesIO(); workbook.save(output); workbook.close(); output.seek(0)
    filename = "Шаблон графика ТО.xlsx"
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})


@router.post("/preview", response_model=ImportPreview)
async def preview_import(file: UploadFile = File(...), user: UserContext = Depends(require_planner), db: Session = Depends(get_db)) -> ImportPreview:
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Поддерживаются файлы XLSX и XLSM")
    content = await file.read()
    if len(content) > 30 * 1024 * 1024:
        raise HTTPException(413, "Файл превышает 30 МБ")
    try:
        preview = ExcelImportService().parse(content, file.filename)
        names = {line_key(line.name) for line in db.scalars(select(ProductionLine).where(ProductionLine.status == "active"))}
        for row in preview.rows:
            if row.line_hint and line_key(row.line_hint) not in names:
                text = f"Неизвестная линия «{row.line_hint}»: линия не будет создана автоматически"
                if preview.template_type in {"production_reference", "capacity_reference", "maintenance_schedule"}:
                    row.errors.append(text)
                    row.valid = False
                else:
                    row.warnings.append(text)
            elif preview.template_type in {"ohl_daily", "quarter_weekly"} and not row.line_hint:
                row.warnings.append("Артикул отсутствует в справочнике мощностей: он будет добавлен в список нераспределённых")
        preview.valid_rows = sum(row.valid for row in preview.rows)
        preview.invalid_rows = len(preview.rows) - preview.valid_rows
        return preview
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc


@router.post("/confirm")
def confirm_import(payload: ImportConfirmRequest, db: Session = Depends(get_db), user: UserContext = Depends(require_planner)) -> dict:
    preview = payload.preview
    if not preview.valid_rows:
        raise HTTPException(422, "В файле нет корректных строк")
    order = ImportedOrder(
        source_name=preview.file_name, status="confirmed", mapping_code=preview.mapping_code,
        template_type=preview.template_type,
        total_rows=preview.total_rows, valid_rows=preview.valid_rows, invalid_rows=preview.invalid_rows,
    )
    db.add(order)
    db.flush()
    db.add(ImportFile(imported_order_id=order.id, original_name=preview.file_name))
    product_by_sku = {product.sku: product for product in db.scalars(select(Product))}
    for row in preview.reference_rows:
        if row.valid and row.sku:
            _upsert_product(db, product_by_sku, row, overwrite=False)

    if preview.template_type == "maintenance_schedule":
        plan = PlanService(db).active_plan()
        if not plan:
            raise HTTPException(422, "Сначала создайте производственный план, затем загрузите график ТО")
        line_by_name = {line_key(line.name): line for line in db.scalars(select(ProductionLine).where(ProductionLine.status == "active"))}
        for existing in plan.schedule_items:
            if existing.source == "maintenance_import" and existing.schedule_kind in {"maintenance", "cleaning", "downtime"}:
                existing.excluded = True
        imported = 0
        valid_rows = [row for row in preview.rows if row.valid and row.line_hint and row.requested_date and row.duration_hours and row.event_kind]
        if valid_rows:
            start = min(row.requested_date for row in valid_rows)
            end = max(row.requested_date for row in valid_rows)
            plan.horizon_start = min(plan.horizon_start, start)
            plan.horizon_end = max(plan.horizon_end, end)
            ensure_line_capacities(db, list(line_by_name.values()), plan.horizon_start, plan.horizon_end)
        for row in valid_rows:
            line = line_by_name.get(line_key(row.line_hint))
            if not line:
                continue
            plan.schedule_items.append(ProductionScheduleItem(
                line_id=line.id, production_date=row.requested_date, shift=row.shift or "day", sequence=0,
                quantity=Decimal("0"), quantity_kg=Decimal("0"), required_hours=row.duration_hours, duration_hours=row.duration_hours,
                schedule_kind=row.event_kind, reason=row.product_name or "Плановое ТО", source_kind="generic",
                source="maintenance_import", locked=True, status=ScheduleStatus.PLANNED,
            ))
            imported += 1
        service = PlanService(db)
        order.status = "maintenance_imported"
        db.flush()
        demands = service._latest_source_demands()
        if demands:
            service.calculate(plan, demands, "maintenance_import")
        else:
            service.refresh_sequence_and_cleanings(plan)
            service.recalculate_load(plan)
            service.create_version(plan, "maintenance_import", f"График ТО: {preview.file_name} · {imported} событий")
        db.add(AuditEvent(username=user.username, action="maintenance_schedule_imported", entity_type="production_plan", entity_id=str(plan.id), details={"file_name": preview.file_name, "events": imported}))
        db.commit()
        return {"order_id": order.id, "plan": plan_dict(db, plan), "maintenance_updated": imported}

    if preview.template_type in {"production_reference", "legacy_reference", "capacity_reference"}:
        updated = 0
        for row in preview.rows:
            if row.sku:
                _upsert_product(db, product_by_sku, row, overwrite=row.valid)
            if not row.valid:
                continue
            product = _upsert_product(db, product_by_sku, row)
            if preview.template_type in {"production_reference", "capacity_reference"}:
                _upsert_capability(db, product, row)
            updated += 1
        order.status = "reference_imported"
        db.add(AuditEvent(
            username=user.username, action="reference_imported", entity_type="imported_order",
            entity_id=str(order.id), details={"file_name": preview.file_name, "template_type": preview.template_type, "updated": updated},
        ))
        db.commit()
        send_notification(
            db, "reference_imported", f"PLAN Portal: загружен справочник {preview.file_name}",
            f"Пользователь {user.display_name} обновил справочник. Обработано строк: {updated}.",
        )
        return {"order_id": order.id, "plan": None, "reference_updated": updated}

    source_kind = {"ohl_daily": "ohl", "quarter_weekly": "zam"}.get(preview.template_type, "generic")
    demands = []
    for row in preview.rows:
        if not row.valid or row.quantity is None or row.requested_date is None or row.due_date is None:
            continue
        product = _upsert_product(db, product_by_sku, row, overwrite=False)
        # Demand workbooks never change line speeds, batches or restrictions.
        if row.line_hint and not any(line_key(line.name) == line_key(row.line_hint) for line in db.scalars(select(ProductionLine).where(ProductionLine.status == "active"))):
            row.warnings = [*row.warnings, f"Неизвестная линия «{row.line_hint}»: проверьте справочник"]
        source_date = row.source_plan_date or row.requested_date
        production_date = source_date - timedelta(days=1) if source_kind == "ohl" and product.advance_status == "АЗ" else source_date
        demand = DemandItem(
            order_id=order.id, product_id=product.id, source_row=row.row_number, sku=row.sku,
            product_name=row.product_name or product.name, quantity=row.quantity,
            source_quantity=row.source_quantity, source_unit=row.source_unit,
            quantity_kg=row.quantity_kg or row.quantity, box_count=row.box_count,
            production_week=row.production_week, exact_date=row.exact_date,
            source_kind=source_kind, source_plan_date=row.source_plan_date or row.requested_date,
            marking_date=row.marking_date, advance_production=row.advance_marking,
            requested_date=production_date if source_kind == "ohl" else row.requested_date,
            due_date=production_date if source_kind == "ohl" else row.due_date, priority=row.priority,
            customer=row.customer, valid=True, validation_errors=row.warnings,
            raw_data={"template_type": preview.template_type, "line_hint": row.line_hint, "advance_marking": row.advance_marking, "source_quantity_exact": str(row.source_quantity), "source_date": str(source_date)},
        )
        db.add(demand)
        demands.append(demand)
    db.flush()
    if not payload.create_plan:
        db.add(AuditEvent(
            username=user.username, action="demand_imported", entity_type="imported_order",
            entity_id=str(order.id), details={"file_name": preview.file_name, "rows": len(demands), "plan_created": False},
        ))
        db.commit()
        return {"order_id": order.id, "plan": None}
    all_demands = _latest_real_demands(db) if payload.merge_into_active and source_kind in {"ohl", "zam"} else demands
    horizon_start = min(item.requested_date for item in all_demands)
    horizon_end = max(item.due_date for item in all_demands)
    _ensure_shift_capacities(db, all_demands, horizon_start, horizon_end)
    plan = PlanService(db).active_plan()
    if not (payload.merge_into_active and plan and plan.name.startswith("Комплексный план ФК")):
        for existing in db.scalars(select(ProductionPlan).where(ProductionPlan.active.is_(True))):
            existing.active = False
        plan = ProductionPlan(name="Комплексный план ФК · ОХЛ + ЗАМ", horizon_start=horizon_start, horizon_end=horizon_end, active=True)
        db.add(plan)
        db.flush()
    else:
        plan.horizon_start = horizon_start
        plan.horizon_end = horizon_end
    PlanService(db).calculate(plan, all_demands, f"import_{source_kind}")
    db.add(AuditEvent(
        username=user.username, action="plan_imported", entity_type="production_plan", entity_id=str(plan.id),
        details={"file_name": preview.file_name, "source_kind": source_kind, "rows": len(demands)},
    ))
    db.commit()
    send_notification(
        db, "plan_imported", f"PLAN Portal: обновлён план {plan.name}",
        f"Пользователь {user.display_name} загрузил файл {preview.file_name}.\nИсточник: {source_kind.upper()}.\nСтрок: {len(demands)}.",
    )
    return {"order_id": order.id, "plan": plan_dict(db, plan)}


def _upsert_product(db: Session, product_by_sku: dict[str, Product], row, overwrite: bool = True) -> Product:
    product = product_by_sku.get(row.sku)
    if not product:
        product = Product(sku=row.sku, name=row.product_name or row.sku)
        db.add(product)
        db.flush()
        product_by_sku[row.sku] = product
    if row.product_name and (overwrite or product.name == product.sku):
        product.name = row.product_name
    for field in ("advance_status", "fk_status"):
        if getattr(product, field) is None and getattr(row, field, None) is not None:
            setattr(product, field, getattr(row, field))
    if overwrite:
        product.unit = row.source_unit or product.unit
    for field in ("unit_weight_kg", "units_per_box", "box_weight_kg", "state", "category", "short_name"):
        value = getattr(row, field, None)
        if value is not None and (overwrite or getattr(product, field) is None):
            setattr(product, field, value)
    for field in ("legacy_quantum_units", "legacy_daily_capacity_units", "legacy_capacity_unit", "reference_source"):
        value = getattr(row, field, None)
        if value is not None:
            setattr(product, field, value)
    if getattr(row, "recipe_component_count", 0):
        product.recipe_component_count = row.recipe_component_count
    return product


def _upsert_capability(db: Session, product: Product, row) -> LineCapability | None:
    if not row.line_hint or not row.speed_kg_hour:
        return None
    line = next((line for line in db.scalars(select(ProductionLine).where(ProductionLine.status == "active")) if line_key(line.name) == line_key(row.line_hint)), None)
    if not line:
        raise HTTPException(422, f"Неизвестная линия «{row.line_hint}». Импорт не создаёт линии: исправьте сопоставление в справочнике.")
    else:
        line.workshop_code, line.workshop_name = workshop_for_line(line.name)
    capability = db.scalar(select(LineCapability).where(
        LineCapability.line_id == line.id, LineCapability.product_id == product.id,
    ))
    if not capability:
        capability = LineCapability(line_id=line.id, product_id=product.id, units_per_hour=row.speed_kg_hour)
        db.add(capability)
        # One source row is created per date, so the same SKU/line pair repeats many
        # times in OHL. Flush now to make the new pair visible to the next lookup.
        db.flush()
    capability.units_per_hour = row.speed_kg_hour
    capability.speed_unit = "кг/час"
    capability.batch_quantum_kg = row.batch_quantum_kg
    capability.min_order_kg = row.min_order_kg
    capability.capacity_type = row.capacity_type
    capability.restrictions = row.restrictions
    capability.min_batch = row.batch_quantum_kg
    if getattr(row, "available_hours", None) is not None or getattr(row, "line_status", None):
        capability.technological_constraints = {
            **(capability.technological_constraints or {}),
            "source_available_hours": str(row.available_hours) if row.available_hours is not None else None,
            "source_line_status": row.line_status,
        }
    return capability


def _ensure_shift_capacities(db: Session, demands: list[DemandItem], start, end) -> None:
    product_ids = {item.product_id for item in demands if item.product_id}
    line_ids = set(db.scalars(select(LineCapability.line_id).where(LineCapability.product_id.in_(product_ids))))
    lines = list(db.scalars(select(ProductionLine).where(ProductionLine.id.in_(line_ids), ProductionLine.status == "active")))
    ensure_line_capacities(db, lines, start, end)


def _latest_real_demands(db: Session) -> list[DemandItem]:
    return [item for item in PlanService(db)._latest_source_demands() if item.valid]
