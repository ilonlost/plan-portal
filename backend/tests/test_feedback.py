from io import BytesIO

from fastapi import Request
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.api.routes.feedback import FeedbackCreate, FeedbackUpdate, create_feedback, export_feedback, list_feedback, update_feedback
from app.core.config import settings
from app.core.security import UserContext
from app.db.base import Base


def _request() -> Request:
    return Request({"type": "http", "method": "POST", "path": "/api/feedback", "headers": [], "client": ("127.0.0.1", 12345)})


def test_feedback_is_private_to_author_and_admin_can_process_it(monkeypatch) -> None:
    monkeypatch.setattr(settings, "email_enabled", False)
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine, expire_on_commit=False)()
    author = UserContext("planner.user", "Анна Планер", "planner", "planner@agrohold.ru")
    admin = UserContext("portal.admin", "Администратор", "admin", "admin@agrohold.ru")

    created = create_feedback(
        FeedbackCreate(category="suggestion", subject="Добавить фильтр", message="Нужен фильтр в недельном плане по сменам."),
        _request(), db, author,
    )

    assert created["ok"] is True
    assert created["entry"]["notification_status"] == "skipped"
    assert list_feedback(status="", query="", db=db, user=author)["entries"][0]["subject"] == "Добавить фильтр"
    assert list_feedback(status="", query="", db=db, user=UserContext("other", "Другой", "viewer"))["entries"] == []

    updated = update_feedback(
        created["entry"]["id"], FeedbackUpdate(status="resolved", comment="Фильтр добавлен в план."), db, admin,
    )

    assert updated["entry"]["status"] == "resolved"
    all_entries = list_feedback(status="", query="", db=db, user=admin)
    assert all_entries["entries"][0]["it_comment"] == "Фильтр добавлен в план."
    assert len(all_entries["events"]) == 3

    private_export = export_feedback(status="", query="", db=db, user=UserContext("other", "Другой", "viewer"))
    private_sheet = load_workbook(BytesIO(private_export.body), data_only=False)["Обратная связь"]
    assert private_sheet.max_row == 1

    admin_export = export_feedback(status="resolved", query="фильтр", db=db, user=admin)
    assert admin_export.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    sheet = load_workbook(BytesIO(admin_export.body), data_only=False)["Обратная связь"]
    assert sheet.freeze_panes == "A2"
    assert sheet["A2"].value == created["entry"]["id"]
    assert sheet["D2"].value == "Добавить фильтр"
    assert sheet["O1"].value == "Правки / комментарий"
    db.close()


def test_feedback_export_escapes_excel_formulas(monkeypatch) -> None:
    monkeypatch.setattr(settings, "email_enabled", False)
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine, expire_on_commit=False)()
    author = UserContext("planner.user", "Анна Планер", "planner", "planner@agrohold.ru")

    create_feedback(
        FeedbackCreate(category="problem", subject="=1+1", message="@SUM(A1:A2) must stay plain text"),
        _request(), db, author,
    )
    sheet = load_workbook(BytesIO(export_feedback(status="", query="", db=db, user=author).body), data_only=False)["Обратная связь"]

    assert sheet["D2"].data_type != "f"
    assert sheet["D2"].value == "'=1+1"
    assert sheet["E2"].data_type != "f"
    assert sheet["E2"].value == "'@SUM(A1:A2) must stay plain text"
    db.close()
