from collections import defaultdict
from datetime import date, time, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import asyncio
import json

import pytest
from fastapi import FastAPI, UploadFile
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select, func
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from sqlalchemy.orm.exc import StaleDataError

from app.db.base import Base
from app.db.session import get_db
from app.core.config import settings
from app.core.security import UserContext, create_session_token, parse_session_token
from app.api.routes import session, plans, imports, catalog, admin, planner_admin, advance_confirmations, production_fact, lines
from app.models.entities import User, Product, ProductionLine, ProductionPlan, ProductionPlanVersion, ProductionScheduleItem, LineCapability, LineCapacity, LineComment, DemandItem, ImportedOrder
from app.schemas.common import ImportConfirmRequest
from app.services.import_service import ExcelImportService
from app.services.plan_service import PlanService, schedule_item_dict
from app.services.planning_engine import PlanningEngine, DemandInput, CapabilityInput, CapacityInput
from app.services.notification_service import build_plan_email_html
from app.seed import WORKSHOPS


@pytest.fixture
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False, autoflush=False) as session_db:
        for code, (name, lines) in WORKSHOPS.items():
            for line in lines:
                session_db.add(ProductionLine(code=line, name=line, workshop_code=code, workshop_name=name, working_hours=22, schedule_code="day_night_daily"))
        session_db.add(User(username="regular.admin", display_name="Administrator", role="admin"))
        session_db.commit()
        yield session_db
    engine.dispose()


@pytest.fixture
def client(db, monkeypatch):
    monkeypatch.setattr(settings, "app_env", "production")
    monkeypatch.setattr(settings, "auth_mode", "ldap")
    monkeypatch.setattr(settings, "session_secret", "test-secret-with-more-than-32-characters")
    app = FastAPI()
    app.include_router(session.router)
    app.include_router(plans.router)
    app.include_router(lines.router)
    app.include_router(imports.router)
    app.include_router(advance_confirmations.router)
    app.include_router(catalog.router)
    app.include_router(admin.router)
    app.include_router(planner_admin.router)
    app.include_router(production_fact.router)
    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)
    client.cookies.set(settings.session_cookie_name, create_session_token(UserContext("regular.admin", "Administrator", "admin")))
    return client


def test_production_fact_is_admin_only_and_exports_detail(client, db, monkeypatch):
    monkeypatch.setattr(settings, "production_fact_database_url", "")
    response = client.get("/production-fact?start=2026-09-01&end=2026-09-02")
    assert response.status_code == 200
    body = response.json()
    assert body["source"] == "csb_dwh" and body["status"] == "not_configured"
    assert len(body["centers"]) == 12
    assert body["total_count"] == 0 and not body["rows"]
    assert "production_hours" not in body and "articles" not in body
    detail = client.get("/production-fact/cost-centers/5810?start=2026-09-01&end=2026-09-02")
    assert detail.status_code == 200 and not detail.json()["rows"]
    assert client.get("/production-fact/export.xlsx?start=2026-09-01&end=2026-09-02").status_code == 503

    stored = db.scalar(select(User).where(User.username == "regular.admin"))
    stored.role = "viewer"
    db.commit()
    assert client.get("/production-fact").status_code == 403
    stored.role = "admin"
    db.commit()
    assert client.get("/production-fact?start=2026-01-01&end=2026-09-01").status_code == 422


def test_production_rejects_demo_and_header_bypass(client, db, monkeypatch):
    client.cookies.clear()
    db.add(User(username="demo.admin", display_name="Legacy demo", role="admin", active=True))
    db.commit()
    for mode in ("ldap", "mock", "typo"):
        monkeypatch.setattr(settings, "auth_mode", mode)
        assert client.post("/session/login", json={"username": "demo.admin", "password": "demo"}).status_code == 401
        assert client.get("/session/me", headers={"X-User": "demo.admin"}).status_code == 401
        assert client.get("/session/mode").json() == {"auth_mode": "ldap"}
    monkeypatch.setattr(settings, "app_env", "development")
    monkeypatch.setattr(settings, "auth_mode", "local")
    token = create_session_token(UserContext("demo.admin", "Demo", "admin"))
    monkeypatch.setattr(settings, "app_env", "production")
    monkeypatch.setattr(settings, "auth_mode", "ldap")
    assert parse_session_token(token) is None
    # Ordinary LDAP sessions remain usable and roles are read from the database.
    token = create_session_token(UserContext("regular.admin", "Administrator", "viewer"))
    client.cookies.set(settings.session_cookie_name, token)
    assert client.get("/session/me").json()["role"] == "admin"
    db.scalar(select(User).where(User.username == "regular.admin")).active = False
    db.commit()
    assert client.get("/session/me").status_code == 403


def make_plan(db, quantity=100):
    day = date(2026, 9, 8)
    line = db.scalar(select(ProductionLine).where(ProductionLine.name == "Сэндвичи"))
    product = Product(sku="101", name="Test", box_weight_kg=Decimal("1"))
    order = ImportedOrder(source_name="ОХЛ.xlsx", template_type="ohl_daily")
    db.add_all([product, order]); db.flush()
    db.add(LineCapability(line_id=line.id, product_id=product.id, units_per_hour=100))
    db.add(LineCapacity(line_id=line.id, capacity_date=day, shift="day", available_hours=1, manual_override=True))
    db.add(LineCapacity(line_id=line.id, capacity_date=day, shift="night", available_hours=0, manual_override=True))
    demand = DemandItem(order_id=order.id, product=product, source_row=3, sku="101", product_name="Test", quantity=quantity, source_kind="ohl", source_plan_date=day, source_quantity=quantity, source_unit="кг", requested_date=day, due_date=day, exact_date=True)
    plan = ProductionPlan(name="Комплексный план ФК · ОХЛ + ЗАМ", horizon_start=day, horizon_end=day)
    db.add_all([demand, plan]); db.flush()
    PlanService(db).calculate(plan, [demand])
    return plan, demand


def test_matrix_keeps_requested_weeks_outside_plan_horizon(client, db):
    plan, _ = make_plan(db)
    db.commit()
    for days in (1, 7, 21):
        response = client.get(f"/plans/active/matrix?start=2026-10-01&days={days}")
        assert response.status_code == 200
        body = response.json()
        assert len(body["dates"]) == days
        assert body["dates"][0] == "2026-10-01"
        assert all(len(line["cells"]) == days for shop in body["workshops"] for line in shop["lines"])


def test_seed_preserves_existing_plan_tasks_and_line_configuration(db, monkeypatch):
    from app import seed as bootstrap
    from sqlalchemy.orm import sessionmaker
    plan, _ = make_plan(db)
    line = db.scalar(select(ProductionLine).where(ProductionLine.name == "Сэндвичи"))
    line.csb_line_code = "CUSTOM-CSB"
    db.commit()
    def snapshot():
        db.expire_all()
        return ([tuple(getattr(row, column.name) for column in ProductionScheduleItem.__table__.columns) for row in db.scalars(select(ProductionScheduleItem).order_by(ProductionScheduleItem.id))], db.get(ProductionPlan, plan.id).revision, db.get(ProductionLine, line.id).csb_line_code)
    before = snapshot()
    monkeypatch.setattr(bootstrap, "engine", db.get_bind())
    monkeypatch.setattr(bootstrap, "SessionLocal", sessionmaker(bind=db.get_bind()))
    monkeypatch.setattr(settings, "app_env", "production")
    monkeypatch.setattr(settings, "auth_mode", "ldap")
    bootstrap.seed()
    bootstrap.seed()
    assert snapshot() == before


def test_user_access_api_is_individual_and_applies_to_existing_session(client, db):
    target = User(username="limited.viewer", display_name="Limited", role="viewer", active=True)
    other = User(username="other.viewer", display_name="Other", role="viewer", active=True)
    db.add_all([target, other]); db.commit()
    response = client.patch(f"/admin/users/{target.id}", json={"role": "viewer", "active": True, "section_permissions": {"catalog": False, "sources": False, "fact": True}})
    assert response.status_code == 200
    assert other.section_permissions == {}
    client.cookies.set(settings.session_cookie_name, create_session_token(UserContext(target.username, target.display_name, "viewer")))
    assert client.get("/catalog").status_code == 403
    assert client.get("/production-fact").status_code == 200
    assert client.get("/session/me").json()["section_visibility"]["catalog"] is False


def test_advance_confirmation_preview_apply_and_fact_export(client, db):
    plan, demand = make_plan(db, quantity=100)
    demand.product.advance_status = "АЗ"
    db.commit()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Материал", "Всего КГ", "Сколько произвели"])
    sheet.append(["101", 75.5, 72])
    content = BytesIO()
    workbook.save(content)

    response = client.post(
        "/advance-confirmations/preview",
        files={"file": ("Потверждение АЗ 08.09.2026.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview["marking_date"] == "2026-09-08"
    assert preview["rows"][0]["production_date"] == "2026-09-07"
    assert preview["summary"]["total_rows"] == 1
    assert preview["rows"][0]["line_name"] == "Сэндвичи"
    assert Decimal(str(preview["rows"][0]["quantity_kg"])) == Decimal("75.5")

    response = client.post("/advance-confirmations/apply-workbook", files={
        "file": (preview["file_name"], content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    }, data={"payload": json.dumps({"file_name": preview["file_name"], "rows": preview["rows"]})})
    assert response.status_code == 200, response.text
    assert response.json()["plan_recalculated"] is True
    db.refresh(demand)
    assert demand.quantity == Decimal("75.500000")
    assert demand.requested_date == date(2026, 9, 7)
    assert demand.raw_data["preferred_line_id"]
    history = client.get("/advance-confirmations").json()
    assert history[0]["file_name"] == preview["file_name"]
    assert history[0]["items"][0]["production_date"] == "2026-09-07"
    assert Decimal(str(history[0]["items"][0]["actual_quantity_kg"])) == Decimal("72")
    exported = client.get(f"/advance-confirmations/{history[0]['id']}/export.xlsx")
    assert exported.status_code == 200
    fact = load_workbook(BytesIO(exported.content), read_only=True, data_only=True)
    assert fact.active["A2"].value == "101"
    assert fact.active["B2"].value == 75.5
    assert fact.active["C2"].value == 72


def test_pc_auto_wash_can_be_removed_without_deleting_production(db):
    day = date(2026, 9, 14)
    line = db.scalar(select(ProductionLine).where(ProductionLine.workshop_code == "PC"))
    first = Product(sku="WASH-A", name="Булка A", mono_group="A")
    second = Product(sku="WASH-B", name="Булка B", mono_group="B")
    plan = ProductionPlan(name="Wash test", horizon_start=day, horizon_end=day)
    db.add_all([first, second, plan]); db.flush()
    plan.schedule_items.extend([
        ProductionScheduleItem(product_id=first.id, line_id=line.id, production_date=day, shift="day", sequence=1, quantity=100, quantity_kg=100, required_hours=1, schedule_kind="production", source="manual"),
        ProductionScheduleItem(product_id=second.id, line_id=line.id, production_date=day, shift="day", sequence=2, quantity=100, quantity_kg=100, required_hours=1, schedule_kind="production", source="manual"),
    ])
    service = PlanService(db); db.flush(); service.refresh_sequence_and_cleanings(plan)
    washes = sorted([item for item in plan.schedule_items if item.schedule_kind == "cleaning" and not item.excluded], key=lambda item: item.sequence)
    assert len(washes) == 2
    service.delete_item(washes[0])
    assert len([item for item in plan.schedule_items if item.schedule_kind == "production" and not item.excluded]) == 2
    assert len([item for item in plan.schedule_items if item.schedule_kind == "cleaning" and not item.excluded]) == 1
    service.calculate(plan, [])
    assert len([item for item in plan.schedule_items if item.schedule_kind == "production" and not item.excluded]) == 2
    assert len([item for item in plan.schedule_items if item.schedule_kind == "cleaning" and not item.excluded]) == 1


def test_manual_priority_reorders_tasks_inside_same_shift(db):
    day = date(2026, 9, 14)
    line = db.scalar(select(ProductionLine).where(ProductionLine.workshop_code == "KC"))
    first_product = Product(sku="ORDER-A", name="Первое задание")
    second_product = Product(sku="ORDER-B", name="Второе задание")
    plan = ProductionPlan(name="Priority test", horizon_start=day, horizon_end=day)
    db.add_all([first_product, second_product, plan]); db.flush()
    db.add_all([
        LineCapability(line_id=line.id, product_id=first_product.id, units_per_hour=100),
        LineCapability(line_id=line.id, product_id=second_product.id, units_per_hour=100),
    ])
    first = ProductionScheduleItem(product_id=first_product.id, line_id=line.id, production_date=day, shift="day", sequence=1, quantity=100, quantity_kg=100, required_hours=1, schedule_kind="production", source="manual")
    second = ProductionScheduleItem(product_id=second_product.id, line_id=line.id, production_date=day, shift="day", sequence=2, quantity=100, quantity_kg=100, required_hours=1, schedule_kind="production", source="manual")
    plan.schedule_items.extend([first, second]); db.flush()
    PlanService(db).update_item(second, {"before_item_id": first.id, "comment": "Поменять приоритет"})
    assert second.sequence < first.sequence
    assert second.sort_rank < first.sort_rank


def test_priority_order_survives_automatic_recalculation(db):
    day = date(2026, 9, 14)
    line = db.scalar(select(ProductionLine).where(ProductionLine.workshop_code == "KC"))
    order = ImportedOrder(source_name="ОХЛ.xlsx", template_type="ohl_daily")
    plan = ProductionPlan(name="Persistent priority", horizon_start=day, horizon_end=day)
    products = [Product(sku=f"PRIORITY-{letter}", name=f"Продукт {letter}") for letter in "ABC"]
    db.add_all([order, plan, *products]); db.flush()
    db.add_all([LineCapability(line_id=line.id, product_id=product.id, units_per_hour=100) for product in products])
    db.add_all([
        LineCapacity(line_id=line.id, capacity_date=day, shift="day", available_hours=3, manual_override=True),
        LineCapacity(line_id=line.id, capacity_date=day, shift="night", available_hours=0, manual_override=True),
    ])
    demands = [DemandItem(order_id=order.id, product=product, source_row=index + 2, sku=product.sku, product_name=product.name, quantity=100, source_kind="ohl", source_quantity=100, source_unit="кг", requested_date=day, due_date=day, exact_date=True) for index, product in enumerate(products)]
    db.add_all(demands); db.flush()
    service = PlanService(db); service.calculate(plan, demands)
    first, second, third = sorted((item for item in plan.schedule_items if item.schedule_kind == "production"), key=lambda item: item.sequence)
    service.update_item(third, {"before_item_id": second.id, "comment": "C перед B"})
    service.calculate(plan, demands)
    ordered_names = [item.product.name for item in sorted((row for row in plan.schedule_items if row.schedule_kind == "production" and not row.excluded), key=lambda row: row.sequence)]
    assert ordered_names == ["Продукт A", "Продукт C", "Продукт B"]


def test_maintenance_template_download_contains_lines_and_validations(client):
    response = client.get("/imports/maintenance-template.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["График ТО", "Инструкция"]
    sheet = workbook["График ТО"]
    assert [cell.value for cell in sheet[1]] == ["Линия", "Дата", "Смена", "Тип события", "Начало", "Окончание", "Причина / работы"]
    assert sum(1 for row in range(2, sheet.max_row + 1) if sheet.cell(row, 1).value) == 15
    assert len(sheet.data_validations.dataValidation) == 2
    workbook.close()


def test_imported_maintenance_reserves_capacity_during_recalculation(db):
    day = date(2026, 9, 14)
    future_day = day + timedelta(days=5)
    line = db.scalar(select(ProductionLine).where(ProductionLine.workshop_code == "KC"))
    product = Product(sku="TO-CAPACITY", name="Продукт после ТО")
    order = ImportedOrder(source_name="ОХЛ.xlsx", template_type="ohl_daily")
    plan = ProductionPlan(name="Maintenance capacity", horizon_start=day, horizon_end=day)
    db.add_all([product, order, plan]); db.flush()
    db.add(LineCapability(line_id=line.id, product_id=product.id, units_per_hour=100))
    db.add_all([
        LineCapacity(line_id=line.id, capacity_date=day, shift="day", available_hours=4, manual_override=True),
        LineCapacity(line_id=line.id, capacity_date=day, shift="night", available_hours=0, manual_override=True),
    ])
    demand = DemandItem(order_id=order.id, product=product, source_row=2, sku=product.sku, product_name=product.name, quantity=400, source_kind="ohl", source_quantity=400, source_unit="кг", requested_date=day, due_date=day, exact_date=True)
    plan.schedule_items.extend([
        ProductionScheduleItem(line_id=line.id, production_date=day, shift="day", sequence=1, quantity=0, quantity_kg=0, required_hours=2, duration_hours=2, schedule_kind="maintenance", source="maintenance_import", locked=True),
        ProductionScheduleItem(line_id=line.id, production_date=future_day, shift="day", sequence=2, quantity=0, quantity_kg=0, required_hours=1, duration_hours=1, schedule_kind="maintenance", source="maintenance_import", locked=True),
    ])
    db.add(demand); db.flush()
    PlanService(db).calculate(plan, [demand])
    scheduled = sum(item.quantity for item in plan.schedule_items if item.schedule_kind == "production" and item.status.value != "unscheduled")
    assert scheduled == Decimal("400")
    assert any(row.source_kind == "ohl" and row.load_percent > 100 for row in plan.schedule_items)
    assert any(item.schedule_kind == "maintenance" and not item.excluded for item in plan.schedule_items)
    assert plan.horizon_end == future_day


def test_kc_downtime_gets_restart_interval(db):
    day = date(2026, 9, 14)
    line = db.scalar(select(ProductionLine).where(ProductionLine.workshop_code == "KC"))
    plan = ProductionPlan(name="CONS test", horizon_start=day, horizon_end=day)
    db.add(plan); db.flush()
    service = PlanService(db)
    service.create_event(plan, {"line_id": line.id, "production_date": day, "shift": "day", "schedule_kind": "downtime", "duration_hours": Decimal("2"), "reason": "Остановка CONS"})
    restart = next(item for item in plan.schedule_items if item.schedule_kind == "restart" and not item.excluded)
    assert restart.required_hours == Decimal("0.5")
    assert restart.sequence > next(item.sequence for item in plan.schedule_items if item.schedule_kind == "downtime")


def test_manual_event_uses_exact_start_and_end_time(db):
    day = date(2026, 9, 14)
    line = db.scalar(select(ProductionLine).where(ProductionLine.workshop_code == "KC"))
    plan = ProductionPlan(name="Timed event", horizon_start=day, horizon_end=day)
    db.add(plan); db.flush()
    PlanService(db).create_event(plan, {
        "line_id": line.id, "production_date": day, "shift": "night", "schedule_kind": "maintenance",
        "start_time": time(23, 30), "end_time": time(1, 0), "reason": "Ночное ТО",
    })
    event = next(item for item in plan.schedule_items if item.schedule_kind == "maintenance")
    assert event.start_time == time(23, 30)
    assert event.end_time == time(1, 0)
    assert event.required_hours == Decimal("1.50")


def test_lasagna_defaults_to_ohl_start_and_bolognese_finish(db):
    day = date(2026, 9, 14)
    line = db.scalar(select(ProductionLine).where(ProductionLine.name == "Лазанья"))
    products = [
        Product(sku="LAS-OHL", name="Лазанья с курицей и грибами охл"),
        Product(sku="LAS-ZAM", name="Лазанья с ветчиной зам"),
        Product(sku="LAS-BOL-OHL", name="Лазанья Болоньезе охл"),
        Product(sku="LAS-BOL-ZAM", name="Лазанья Болоньезе зам"),
    ]
    plan = ProductionPlan(name="Lasagna order", horizon_start=day, horizon_end=day)
    db.add_all([plan, *products]); db.flush()
    for product, source_kind in zip(products, ("ohl", "zam", "ohl", "zam")):
        plan.schedule_items.append(ProductionScheduleItem(product_id=product.id, line_id=line.id, production_date=day, shift="day", sequence=len(plan.schedule_items) + 1, quantity=100, quantity_kg=100, required_hours=1, schedule_kind="production", source_kind=source_kind, source="manual"))
    PlanService(db).refresh_sequence_and_cleanings(plan)
    ordered = [item.product.sku for item in sorted((row for row in plan.schedule_items if row.schedule_kind == "production"), key=lambda row: row.sequence)]
    assert ordered == ["LAS-OHL", "LAS-BOL-OHL", "LAS-ZAM", "LAS-BOL-ZAM"]
    assert next(item for item in plan.schedule_items if item.schedule_kind == "startup").required_hours == Decimal("1")


def test_stale_edit_and_missing_version_do_not_write(client, db):
    plan, _ = make_plan(db)
    item = next(i for i in plan.schedule_items if i.schedule_kind == "production")
    version = plan.revision
    path = f"/plans/{plan.id}/items/{item.id}"
    assert client.patch(path, json={"quantity": 90}).status_code == 428
    first = client.patch(path, json={"quantity": 90}, headers={"If-Match": str(version)})
    assert first.status_code == 200, first.text
    assert client.patch(path, json={"quantity": 80}, headers={"If-Match": str(version)}).status_code == 409
    assert item.quantity == 90


def test_database_revision_prevents_simultaneous_writers(db):
    plan, _ = make_plan(db)
    with Session(db.bind) as other:
        stale = other.get(ProductionPlan, plan.id)
        plan.name = "First writer"; db.commit()
        stale.name = "Second writer"
        with pytest.raises(StaleDataError):
            other.commit()


def test_product_soft_delete_keeps_plan_and_catalog_relations(client, db):
    plan, demand = make_plan(db)
    item_ids = [item.id for item in plan.schedule_items]
    product_id = demand.product_id
    response = client.delete(f"/catalog/products/{product_id}")
    assert response.status_code == 200, response.text
    assert db.get(Product, product_id).catalog_status == "deleted"
    assert db.get(Product, product_id).active is False
    assert db.get(ProductionPlan, plan.id) is not None
    assert [item.id for item in db.scalars(select(ProductionScheduleItem).where(ProductionScheduleItem.plan_id == plan.id))] == item_ids
    assert client.get("/catalog").json()["products"] == []


def test_catalog_create_restores_soft_deleted_sku_instead_of_duplicate_error(client, db):
    plan, demand = make_plan(db)
    product_id = demand.product_id
    deleted = client.delete(f"/catalog/products/{product_id}")
    assert deleted.status_code == 200, deleted.text

    restored = client.post("/catalog/products", json={
        "sku": demand.sku, "product_name": "Восстановленный товар", "state": "ОХЛ",
        "line_id": db.scalar(select(LineCapability).where(LineCapability.product_id == product_id)).line_id,
        "units_per_hour": 100,
    })
    assert restored.status_code == 200, restored.text
    assert restored.json()["restored"] is True
    assert restored.json()["product_id"] == product_id
    catalog_body = client.get("/catalog").json()
    found = next(row for row in catalog_body["products"] if row["sku"] == demand.sku)
    assert found["product_name"] == "Восстановленный товар"
    assert found["capability_count"] == 1
    assert not catalog_body["unmapped_products"]
    assert db.get(ProductionPlan, plan.id) is not None


def test_catalog_lookup_finds_deleted_sku_with_saved_line_and_parameters(client, db):
    plan, demand = make_plan(db)
    product_id = demand.product_id
    capability = db.scalar(select(LineCapability).where(LineCapability.product_id == product_id))
    assert client.delete(f"/catalog/products/{product_id}").status_code == 200
    assert client.get("/catalog").json()["products"] == []

    response = client.get("/catalog/products/lookup", params={"sku": f"  {demand.sku}  "})
    assert response.status_code == 200, response.text
    product = response.json()
    assert product["product_id"] == product_id
    assert product["catalog_status"] == "deleted"
    assert not product["active"]
    assert product["capabilities"][0]["capability_id"] == capability.id
    assert Decimal(str(product["capabilities"][0]["speed_kg_hour"])) == capability.units_per_hour
    assert client.get("/catalog").json()["products"] == []
    assert db.get(ProductionPlan, plan.id) is not None


def test_catalog_lookup_finds_blocked_unmapped_sku_and_resolves_duplicate(client, db):
    product = Product(sku="Mixed-SKU", name="Без линии", active=False, catalog_status="blocked")
    db.add(product)
    db.commit()
    payload = {"sku": "mixed-sku", "product_name": "Дубликат"}
    assert client.post("/catalog/products", json=payload).status_code == 409
    response = client.get("/catalog/products/lookup", params={"sku": " mixed-sku "})
    assert response.status_code == 200, response.text
    assert response.json()["product_id"] == product.id
    assert response.json()["catalog_status"] == "blocked"
    assert response.json()["capabilities"] == []
    assert client.get("/catalog/products/lookup", params={"sku": "not-in-catalog"}).status_code == 404
    assert db.scalar(select(func.count(Product.id))) == 1


@pytest.mark.parametrize("deleted", [False, True])
def test_catalog_can_create_or_restore_unmapped_article_with_empty_parameters(client, db, deleted):
    if deleted:
        db.add(Product(sku="unmapped", name="Сохранённый", active=False, catalog_status="deleted"))
        db.commit()
    response = client.post("/catalog/products", json={
        "sku": "unmapped", "product_name": "Без линии", "batch_quantum_kg": None,
        "min_order_kg": None, "restrictions": None,
    })
    assert response.status_code == 200, response.text
    assert response.json().get("restored", False) is deleted
    product = client.get("/catalog/products/lookup", params={"sku": "unmapped"}).json()
    assert product["catalog_status"] == "active"
    assert product["capabilities"] == []
    assert db.scalar(select(func.count(Product.id))) == 1


def test_catalog_xlsx_roundtrip_updates_without_deleting_missing_rows(client, db):
    plan, demand = make_plan(db)
    exported = client.get("/catalog/export.xlsx")
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content))
    sheet = workbook["Артикулы"]
    sheet.cell(2, 3).value = "Обновлённое название"
    sheet.cell(2, 4).value = "ЗАМ"
    sheet.cell(2, 10).value = "blocked"
    payload = BytesIO(); workbook.save(payload)
    response = client.post("/catalog/import.xlsx", files={"file": ("catalog.xlsx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 200, response.text
    product = db.get(Product, demand.product_id)
    assert product.name == "Обновлённое название"
    assert product.state == "ЗАМ"
    assert product.catalog_status == "blocked"
    assert db.get(ProductionPlan, plan.id) is not None


def test_bom_request_uses_sku_and_marks_1000_unit_basis(client, db, monkeypatch):
    _, demand = make_plan(db)
    requested = []

    class FakeResponse:
        def raise_for_status(self): pass
        def json(self): return {"components": [{"material": "Тесто", "quantity": 250}, {"material": "Соус", "quantity": 20}]}

    class FakeClient:
        def __init__(self, **kwargs): pass
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def get(self, url, **kwargs): requested.append(url); return FakeResponse()

    monkeypatch.setattr(catalog.httpx, "Client", FakeClient)
    response = client.get(f"/catalog/products/{demand.product_id}/bom")
    assert response.status_code == 200, response.text
    body = response.json()
    assert requested[0].endswith("/101")
    assert body["basis_units"] == 1000
    assert body["columns"] == ["material", "quantity"]
    assert len(body["rows"]) == 2


def test_admin_can_delete_other_user_without_deleting_plan_history(client, db):
    target = User(username="former.user", display_name="Former", role="planner", active=True)
    db.add(target); db.flush()
    day = date(2026, 9, 8)
    plan = ProductionPlan(name="Retained", horizon_start=day, horizon_end=day, created_by_id=target.id)
    db.add(plan); db.flush()
    version = ProductionPlanVersion(plan_id=plan.id, version_number=1, change_type="test", snapshot={}, changed_by_id=target.id)
    db.add(version); db.commit()
    response = client.delete(f"/admin/users/{target.id}")
    assert response.status_code == 200, response.text
    assert db.get(User, target.id) is None
    assert db.get(ProductionPlan, plan.id).created_by_id is None
    assert db.get(ProductionPlanVersion, version.id).changed_by_id is None
    regular = db.scalar(select(User).where(User.username == "regular.admin"))
    assert client.delete(f"/admin/users/{regular.id}").status_code == 409


def test_delete_split_stays_deleted_and_other_volume_survives(db):
    plan, demand = make_plan(db, 150)
    service = PlanService(db)
    demand.source_kind = "zam"; demand.exact_date = False; demand.order.template_type = "quarter_weekly"
    service.calculate(plan, [demand])
    original = next(i for i in plan.schedule_items if i.production_date)
    service.delete_item(original)
    service.calculate(plan, [demand])
    live = [i for i in plan.schedule_items if not i.excluded and i.schedule_kind == "production"]
    assert sum(i.quantity for i in live) == 50
    assert original.excluded
    # Reimport under a different filename and row number uses the same business identity.
    new_order = ImportedOrder(source_name="Новая неделя.xlsx", template_type="quarter_weekly")
    db.add(new_order); db.flush()
    fresh = DemandItem(order_id=new_order.id, product=demand.product, source_row=9, sku=demand.sku, product_name=demand.product_name, quantity=150, source_kind="zam", source_plan_date=demand.source_plan_date, requested_date=demand.requested_date, due_date=demand.due_date, exact_date=False)
    db.add(fresh); db.flush()
    assert [d.id for d in service._latest_source_demands()] == [fresh.id]
    service.calculate(plan, [fresh])
    assert sum(i.quantity for i in plan.schedule_items if not i.excluded and i.schedule_kind == "production") == 50


def test_zam_residual_and_increased_capacity():
    day = date(2026, 9, 8)
    demands = [DemandInput(1, 1, "OHL", Decimal("28.4"), day, day, source_kind="ohl", exact_date=True, box_quantum_kg=Decimal("0.6")), DemandInput(2, 2, "ZAM", Decimal("90"), day, day, source_kind="zam")]
    capabilities = [CapabilityInput(1, 1, Decimal("100")), CapabilityInput(1, 2, Decimal("100"))]
    def run(hours):
        result = PlanningEngine().plan(demands, capabilities, [CapacityInput(1, day, Decimal(hours))], day)
        assert sum(i.required_hours for i in result if i.production_date) <= Decimal(hours)
        assert result[0].quantity == Decimal("28.8")
        return result
    assert any(i.status == "unscheduled" for i in run("1"))
    assert not any(i.status == "unscheduled" for i in run("2"))


def test_catalog_speed_recalculates_without_changing_shift_capacity(client, db):
    plan, demand = make_plan(db, 150)
    capability = db.scalar(select(LineCapability).where(LineCapability.product_id == demand.product_id))
    before = {slot.id: slot.available_hours for slot in db.scalars(select(LineCapacity))}
    response = client.patch(f"/catalog/capabilities/{capability.id}", json={"units_per_hour": 200})
    assert response.status_code == 200, response.text
    assert {slot.id: slot.available_hours for slot in db.scalars(select(LineCapacity))} == before
    assert not any(i.status.value == "unscheduled" for i in plan.schedule_items if not i.excluded)


def test_catalog_quantum_change_updates_manually_moved_plan_task(client, db):
    plan, demand = make_plan(db, 103)
    product = demand.product
    product.box_weight_kg = None
    product.unit_weight_kg = None
    capability = db.scalar(select(LineCapability).where(LineCapability.product_id == product.id))
    capability.batch_quantum_kg = Decimal("10")
    db.commit()
    item = next(row for row in plan.schedule_items if row.schedule_kind == "production")
    service = PlanService(db)
    service.update_item(item, {"quantity": 103, "production_date": item.production_date, "line_id": item.line_id, "locked": True})
    assert item.source == "manual" and item.quantity == Decimal("110")

    response = client.patch(f"/catalog/capabilities/{capability.id}", json={"batch_quantum_kg": 25})
    assert response.status_code == 200, response.text
    assert item.quantity == Decimal("125")
    assert item.batch_count == Decimal("5.000")


def test_catalog_advance_status_change_updates_existing_manual_task_date(client, db):
    plan, demand = make_plan(db)
    demand.product.advance_status = "По графику"
    PlanService(db).calculate(plan, [demand])
    item = next(row for row in plan.schedule_items if row.schedule_kind == "production")
    old_date = demand.requested_date
    assert item.production_date == old_date
    planner_date = old_date + timedelta(days=2)
    PlanService(db).update_item(item, {"production_date": planner_date, "line_id": item.line_id, "locked": True})

    response = client.patch(f"/catalog/products/{demand.product_id}", json={"advance_status": "АЗ"})
    assert response.status_code == 200, response.text
    assert demand.requested_date == old_date - timedelta(days=1)
    assert item.production_date == planner_date - timedelta(days=1)


def test_catalog_line_change_moves_manual_task_to_valid_mapping(client, db):
    plan, demand = make_plan(db)
    capability = db.scalar(select(LineCapability).where(LineCapability.product_id == demand.product_id))
    old_line_id = capability.line_id
    target = db.scalar(select(ProductionLine).where(
        ProductionLine.workshop_code == capability.line.workshop_code,
        ProductionLine.id != old_line_id,
    ))
    item = next(row for row in plan.schedule_items if row.schedule_kind == "production")
    PlanService(db).update_item(item, {"production_date": item.production_date, "line_id": item.line_id, "locked": True})

    response = client.patch(f"/catalog/products/{demand.product_id}", json={
        "capability_id": capability.id, "line_id": target.id, "units_per_hour": 100,
    })
    assert response.status_code == 200, response.text
    assert item.source == "manual"
    assert item.line_id == target.id


def test_product_line_change_immediately_rebuilds_active_plan(client, db):
    plan, demand = make_plan(db)
    capability = db.scalar(select(LineCapability).where(LineCapability.product_id == demand.product_id))
    target = db.scalar(select(ProductionLine).where(ProductionLine.workshop_code == "KC", ProductionLine.id != capability.line_id))
    previous_revision = plan.revision
    response = client.patch(f"/catalog/products/{demand.product_id}", json={
        "capability_id": capability.id, "line_id": target.id, "units_per_hour": 100,
    })
    assert response.status_code == 200, response.text
    assert response.json()["plan_recalculated"] is True
    assert plan.revision > previous_revision
    assert {item.line_id for item in plan.schedule_items if not item.excluded and item.schedule_kind == "production"} == {target.id}


def test_unplaced_refresh_recalculates_and_clears_resolved_catalog_warning(client, db):
    plan, demand = make_plan(db)
    capability = db.scalar(select(LineCapability).where(LineCapability.product_id == demand.product_id))
    db.delete(capability)
    demand.validation_errors = [
        "Артикул отсутствует в справочнике мощностей: он будет добавлен в список нераспределённых"
    ]
    PlanService(db).calculate(plan, [demand])
    assert sum(
        item.status.value == "unscheduled" and not item.excluded
        for item in plan.schedule_items
    ) == 1

    response = client.post("/plans/active/recalculate")
    assert response.status_code == 200, response.text
    assert response.json()["unscheduled_count"] == 1

    line = db.scalar(select(ProductionLine).where(ProductionLine.name == "Сэндвичи"))
    db.add(LineCapability(line_id=line.id, product_id=demand.product_id, units_per_hour=100))
    db.commit()
    response = client.post("/plans/active/recalculate")
    assert response.status_code == 200, response.text
    assert response.json()["unscheduled_count"] == 0
    assert not demand.validation_errors
    assert not any(
        item.status.value == "unscheduled" and not item.excluded
        for item in plan.schedule_items
    )


def test_planner_only_can_recalculate_active_plan(client, db):
    make_plan(db)
    db.scalar(select(User).where(User.username == "regular.admin")).role = "viewer"
    db.commit()
    assert client.post("/plans/active/recalculate").status_code == 403


@pytest.mark.parametrize("role", ["admin", "planner", "master", "viewer"])
def test_paused_downtime_insights_are_hidden_for_every_role(client, db, monkeypatch, role):
    db.scalar(select(User).where(User.username == "regular.admin")).role = role
    line = db.scalar(select(ProductionLine))
    db.add(LineComment(line_id=line.id, comment_date=date(2026, 9, 8), text="Комментарий сохранён",
                       author_username="regular.admin", author_name="Administrator"))
    db.commit()
    monkeypatch.setattr(lines, "load_downtimes", lambda *_: pytest.fail("paused downtime source must not be queried"))
    response = client.get("/lines/insights?start=2026-09-08&end=2026-09-08")
    assert response.status_code == 200
    assert response.json()["source_status"] == "hidden"
    assert response.json()["downtimes"] == [] and response.json()["issues"] == []
    assert response.json()["comments"][0]["text"] == "Комментарий сохранён"


def test_shipping_date_stays_fixed_when_production_item_moves(client, db):
    plan, demand = make_plan(db)
    demand.product.advance_status = "По графику"
    demand.marking_date = date(2026, 9, 8)
    PlanService(db).calculate(plan, [demand])
    item = next(row for row in plan.schedule_items if row.schedule_kind == "production")
    marking_date = item.marking_date
    shipping_date = marking_date - timedelta(days=1)
    next_day = item.production_date + timedelta(days=1)
    client.patch(
        f"/plans/{plan.id}/items/{item.id}",
        headers={"If-Match": str(plan.revision)},
        json={"production_date": next_day.isoformat(), "line_id": item.line_id, "locked": True},
    )
    assert item.production_date == next_day
    assert item.marking_date == marking_date
    from app.services.plan_service import schedule_item_dict
    assert schedule_item_dict(item)["shipping_date"] == shipping_date


def test_advance_date_and_production_day_do_not_shift_twice(db):
    plan, demand = make_plan(db)
    demand.marking_date = demand.source_plan_date
    demand.product.advance_status = "АЗ"
    line = next(i.line for i in plan.schedule_items if i.line)
    line.production_day_start_hour = 15
    PlanService(db).calculate(plan, [demand])
    assert demand.requested_date == date(2026, 9, 7)
    assert demand.marking_date == date(2026, 9, 8)
    demand.product.advance_status = "По графику"
    PlanService(db).calculate(plan, [demand])
    assert demand.requested_date == date(2026, 9, 8)


def test_shifted_headers_week_and_invalid_amount():
    book = Workbook(); sheet = book.active; sheet.title = "План ЗАМ"
    sheet.append(["2027", None, None, None]); sheet.append([])
    sheet.append(["Название", "47w", "Код товара SAP", "Комментарий"])
    sheet.append(["Продукция", 42, 101, "x"])
    sheet.append(["Ошибка", "#REF!", 102, "x"])
    stream = BytesIO(); book.save(stream)
    result = ExcelImportService().parse(stream.getvalue(), "ЗАМ.xlsx")
    assert result.rows[0].requested_date == date.fromisocalendar(2027, 47, 1)
    assert result.rows[0].quantity == 42
    assert result.invalid_rows == 1


def test_manual_quantity_override_is_not_recreated_on_recalculation(db):
    plan, demand = make_plan(db)
    service = PlanService(db)
    item = next(i for i in plan.schedule_items if i.schedule_kind == "production")
    service.update_item(item, {"quantity": 90})
    service.calculate(plan, [demand])
    assert sum(i.quantity for i in plan.schedule_items if not i.excluded and i.schedule_kind == "production") == 90


def test_production_and_marking_dates_can_be_edited_independently(client, db):
    plan, _ = make_plan(db)
    item = next(i for i in plan.schedule_items if i.schedule_kind == "production")
    original_production_date = item.production_date
    marking_date = original_production_date + timedelta(days=2)
    response = client.patch(
        f"/plans/{plan.id}/items/{item.id}",
        json={"marking_date": marking_date.isoformat()},
        headers={"If-Match": str(plan.revision)},
    )
    assert response.status_code == 200, response.text
    assert item.production_date == original_production_date
    assert item.marking_date == marking_date

    response = client.patch(
        f"/plans/{plan.id}/items/{item.id}",
        json={"production_date": (original_production_date + timedelta(days=1)).isoformat()},
        headers={"If-Match": str(plan.revision)},
    )
    assert response.status_code == 200, response.text
    assert item.production_date == original_production_date + timedelta(days=1)
    assert item.marking_date == marking_date


def test_unknown_reference_line_is_rejected_without_inserting_it(client, db):
    product = Product(sku="unknown-line", name="Test"); db.add(product); db.commit()
    from app.schemas.common import ImportRow
    with pytest.raises(Exception) as caught:
        imports._upsert_capability(db, product, ImportRow(row_number=1, sku=product.sku, line_hint="Новая случайная линия", speed_kg_hour=100))
    assert caught.value.status_code == 422
    assert db.scalar(select(func.count(ProductionLine.id))) == 15


def test_department_templates_do_not_leak_production_fields():
    day = date(2026, 9, 8)
    items = [{"sku": "101", "product_name": "<Test>", "production_date": day, "marking_date": day + timedelta(days=1), "line_name": "PRIVATE-LINE", "required_hours": 2, "quantity_kg": 40, "box_count": 10, "shift": "night", "status": "planned"}]
    warehouse = build_plan_email_html({}, "Plan", day, day, items, "warehouse")
    materials = build_plan_email_html({}, "Plan", day, day, items, "materials")
    production = build_plan_email_html({}, "Plan", day, day, items)
    assert "PRIVATE-LINE" not in warehouse and "08.09.2026" in warehouse
    assert "&lt;Test&gt;" in warehouse
    assert "не заявка на сырьё" in materials
    assert "PRIVATE-LINE" in production


@pytest.mark.parametrize("name,kind", [("ОХЛ.xlsx", "ohl_daily"), ("ЗАМ.xlsx", "quarter_weekly"), ("Мощности для заливки на портал!!!!!!! 08.09.2026.xlsm", "production_reference")])
def test_supplied_workbooks_reconcile(name, kind):
    path = Path.home() / "Downloads" / name
    if not path.exists():
        pytest.skip("Real workbook is local and intentionally not committed")
    preview = ExcelImportService().parse(path.read_bytes(), name)
    assert preview.template_type == kind
    workbook = load_workbook(path, read_only=True, data_only=True)
    if kind == "ohl_daily":
        sheet = workbook["ОХЛ"]
        headers = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
        cols = [i for i, value in enumerate(headers) if ExcelImportService._header_date(value, 2026)]
        expected = [(r, i, Decimal(str(values[i]))) for r, values in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3) for i in cols if isinstance(values[i], (float, int)) and values[i] > 0 and values[4] and values[5]]
        assert len(preview.rows) == len(expected) == 802
        assert sum(row.source_quantity for row in preview.rows) == sum(v for _, _, v in expected)
        assert {row.production_week for row in preview.rows} == {36, 37, 38}
        assert preview.reference_rows
    elif kind == "quarter_weekly":
        sheet = workbook["План ЗАМ"]
        expected = [Decimal(str(values[i])) for values in sheet.iter_rows(min_row=3, values_only=True) for i in range(4, 8) if isinstance(values[i], (float, int)) and values[i] > 0 and values[1] and values[2]]
        assert len(preview.rows) == len(expected) == 129
        assert preview.valid_rows == 127
        assert preview.invalid_rows == 2
        assert sum(row.source_quantity for row in preview.rows) == sum(expected)
        assert {row.production_week for row in preview.rows} == {36, 37, 38, 39}
    else:
        first = preview.rows[0]
        assert first.speed_kg_hour == Decimal("288.8")
        assert first.batch_quantum_kg == Decimal("146.52888847267886")
        assert first.min_order_kg == Decimal("439.58666541803655")
    workbook.close()


def test_real_import_pipeline_preserves_capacity_and_status(db, monkeypatch):
    root = Path.home() / "Downloads"
    names = ["Мощности для заливки на портал!!!!!!! 08.09.2026.xlsm", "ОХЛ.xlsx", "ЗАМ.xlsx"]
    if not all((root / name).exists() for name in names):
        pytest.skip("Real workbooks not present")
    monkeypatch.setattr(imports, "send_notification", lambda *args, **kwargs: None)
    user = UserContext("regular.admin", "Administrator", "admin")
    snapshots = []
    for name in names:
        preview = asyncio.run(imports.preview_import(UploadFile(filename=name, file=BytesIO((root / name).read_bytes())), user, db))
        result = imports.confirm_import(ImportConfirmRequest(preview=preview), db, user)
        snapshots.append({c.id: (c.units_per_hour, c.batch_quantum_kg, c.min_order_kg) for c in db.scalars(select(LineCapability))})
    assert snapshots[0] == snapshots[1] == snapshots[2]
    assert db.scalar(select(func.count(ProductionLine.id))) == 15
    service = PlanService(db)
    plan = service.active_plan()
    assert plan and any(item.status.value == "unscheduled" for item in plan.schedule_items)
    usage = defaultdict(lambda: Decimal("0"))
    for item in plan.schedule_items:
        if item.production_date and not item.excluded:
            usage[(item.line_id, item.production_date, item.shift)] += item.required_hours
    capacity = {(c.line_id, c.capacity_date, c.shift): c.available_hours if c.available else 0 for c in db.scalars(select(LineCapacity))}
    for key, hours in usage.items():
        if hours > capacity.get(key, 0):
            assert any(row.source_kind == "ohl" and (row.line_id, row.production_date, row.shift) == key and any("Загрузка линии" in warning for warning in row.warnings) for row in plan.schedule_items)
    expected = len([item for item in plan.schedule_items if not item.excluded])
    service.calculate(plan, service._latest_source_demands())
    assert len([item for item in plan.schedule_items if not item.excluded]) == expected
    product = db.scalar(select(Product).where(Product.advance_status == "АЗ"))
    assert product
    product.advance_status = "По графику"; product.fk_status = "Изменено вручную"; db.commit()
    preview = ExcelImportService().parse((root / "ОХЛ.xlsx").read_bytes(), "Повтор ОХЛ.xlsx")
    imports.confirm_import(ImportConfirmRequest(preview=preview), db, user)
    assert product.advance_status == "По графику"
    assert product.fk_status == "Изменено вручную"
    assert len(service._latest_source_demands()) == 802 + 127


def test_planner_administration_restricts_roles_and_hides_integration_settings(client, db):
    actor = db.scalar(select(User).where(User.username == "regular.admin"))
    actor.role = "planner"
    protected = User(username="protected.admin", display_name="Protected", role="admin")
    target = User(username="colleague", display_name="Colleague", role="viewer")
    db.add_all([protected, target]); db.commit()
    body = client.get("/planner-admin").json()
    assert {row["role"] for row in body["users"]} <= {"planner", "viewer"}
    assert "smtp" not in json.dumps(body).lower()
    for role in ("admin", "master"):
        assert client.patch(f"/planner-admin/users/{target.id}", json={"role": role}).status_code == 422
        assert client.post("/planner-admin/users", json={"username": "new", "role": role}).status_code == 422
    assert client.patch(f"/planner-admin/users/{protected.id}", json={"role": "viewer"}).status_code == 403
    assert client.patch(f"/planner-admin/users/{actor.id}", json={"role": "viewer"}).status_code == 409
    assert client.patch(f"/planner-admin/users/{target.id}", json={"role": "planner", "section_permissions": {"sources": False}}).status_code == 200
    assert db.get(User, target.id).role == "planner"
    assert db.get(User, target.id).section_permissions["sources"] is False
    assert client.post("/planner-admin/users", json={"username": "new.viewer", "role": "viewer"}).status_code == 200
    assert client.put("/admin/mail-configuration", json={"configuration": {}}).status_code == 403
    assert client.get("/admin/overview").status_code == 403
    actor.role = "viewer"; db.commit()
    assert client.get("/planner-admin").status_code == 403
    assert client.delete("/planner-admin/sources/1").status_code == 403
    assert client.put("/planner-admin/recipients", json={"production": "", "materials": "", "warehouse": ""}).status_code == 403
    actor.role = "master"; db.commit()
    assert client.get("/planner-admin").status_code == 403


def test_planner_recipients_are_separate_and_exact_for_preview_and_delivery(client, db, monkeypatch):
    from app.services.settings_service import get_mail_configuration, save_mail_configuration
    from app.services.notification_service import send_notification, resolve_plan_recipients
    plan, _ = make_plan(db)
    db.scalar(select(User).where(User.username == "regular.admin")).role = "planner"
    configuration = get_mail_configuration(db)
    configuration.update({"smtp_host": "private.mail", "notification_emails": "legacy@example.com", "event_recipients": {"planning": "planning-events@example.com"}, "enabled": False})
    save_mail_configuration(db, configuration, "test"); db.commit()
    recipients = {"production": "Production@example.com;production@example.com", "warehouse": "warehouse@example.com", "materials": "raw@example.com"}
    assert client.put("/planner-admin/recipients", json={**recipients, "smtp_host": "evil"}).status_code == 422
    assert client.put("/planner-admin/recipients", json={**recipients, "materials": "invalid"}).status_code == 422
    assert client.put("/planner-admin/recipients", json=recipients).status_code == 200
    assert get_mail_configuration(db)["smtp_host"] == "private.mail"
    for audience, expected in (("production", "production@example.com"), ("warehouse", "warehouse@example.com"), ("materials", "raw@example.com")):
        preview = client.get(f"/admin/mail-preview?start=2026-09-08&end=2026-09-08&audience={audience}")
        assert preview.status_code == 200
        assert preview.json()["recipients"] == [expected]
        result = client.post(f"/plans/{plan.id}/email", json={"start": "2026-09-08", "end": "2026-09-08", "audience": audience, "recipients": ["extra@example.com"]})
        assert result.status_code == 200
        assert result.json()["recipients"] == [expected, "extra@example.com"]
    assert resolve_plan_recipients(db, "warehouse", ["master@example.com"]) == ["warehouse@example.com"]
    log = send_notification(db, "schedule_item_updated", "test", "test", ["exact@example.com"], exact_recipients=True)
    assert log.recipients == ["exact@example.com"]


def test_planner_line_downtime_settings_are_per_line_and_stay_hidden(client, db):
    db.scalar(select(User).where(User.username == "regular.admin")).role = "planner"; db.commit()
    first, second = list(db.scalars(select(ProductionLine).limit(2)))
    assert client.put(f"/planner-admin/lines/{first.id}/downtime", json={"partial_percent": 35, "full_percent": 90}).status_code == 200
    for value in (-1, 101, 20.5):
        assert client.put(f"/planner-admin/lines/{first.id}/downtime", json={"partial_percent": value, "full_percent": 80}).status_code == 422
    body = client.get("/planner-admin").json()
    rows = {row["id"]: row for row in body["lines"]}
    assert (rows[first.id]["partial_percent"], rows[first.id]["full_percent"]) == (35, 90)
    assert (rows[second.id]["partial_percent"], rows[second.id]["full_percent"]) == (50, 80)
    assert client.get("/lines/insights?start=2026-09-08&end=2026-09-08").json()["source_status"] == "hidden"


def test_removing_latest_source_removes_its_tasks_without_resurrecting_older_upload(client, db):
    plan, demand = make_plan(db)
    older = ImportedOrder(source_name="Old.xlsx", template_type="ohl_daily", imported_at=demand.order.imported_at - timedelta(days=1))
    db.add(older); db.commit()
    db.scalar(select(User).where(User.username == "regular.admin")).role = "planner"; db.commit()
    source_id = demand.order_id
    assert plan.schedule_items
    response = client.delete(f"/planner-admin/sources/{source_id}")
    assert response.status_code == 200, response.text
    assert response.json()["removed_tasks"] > 0
    assert db.get(ImportedOrder, source_id).status == "deleted"
    assert not PlanService(db)._latest_source_demands()
    assert not any(item.demand_item_id == demand.id for item in plan.schedule_items)
    assert all(row["id"] != source_id for row in client.get("/planner-admin").json()["sources"])
    assert all(row["id"] != source_id for row in client.get("/catalog").json()["sources"])
    assert client.delete(f"/planner-admin/sources/{source_id}").status_code == 404
    assert db.get(Product, demand.product_id) is not None


def test_archiving_reference_does_not_rebuild_plan_or_undo_catalogue(client, db):
    plan, demand = make_plan(db)
    source = ImportedOrder(source_name="Reference.xlsx", template_type="production_reference", status="reference_imported")
    db.add(source); db.commit()
    before = [(item.id, item.quantity) for item in plan.schedule_items]
    response = client.delete(f"/planner-admin/sources/{source.id}")
    assert response.status_code == 200
    assert response.json()["plan_recalculated"] is False
    assert [(item.id, item.quantity) for item in plan.schedule_items] == before
    assert db.get(Product, demand.product_id).name == "Test"


def test_removing_older_source_keeps_latest_demand_and_unrelated_manual_event(client, db):
    plan, demand = make_plan(db)
    older = ImportedOrder(source_name="Old.xlsx", template_type="ohl_daily", imported_at=demand.order.imported_at - timedelta(days=1))
    manual = ProductionScheduleItem(plan=plan, production_date=plan.horizon_start, line_id=plan.schedule_items[0].line_id, source="manual", schedule_kind="gas", quantity=0, required_hours=Decimal("0.1"), duration_hours=Decimal("0.1"), reason="Газ")
    db.add_all([older, manual]); db.commit()
    manual_id = manual.id
    response = client.delete(f"/planner-admin/sources/{older.id}")
    assert response.status_code == 200
    assert [row.id for row in PlanService(db)._latest_source_demands()] == [demand.id]
    assert db.get(ProductionScheduleItem, manual_id).reason == "Газ"
    assert any(item.demand_item_id == demand.id for item in plan.schedule_items)


def test_line_conversation_appends_replies_and_keeps_legacy_comment(client, db):
    from app.models.entities import LineCommentMessage
    line = db.scalar(select(ProductionLine))
    db.add(LineComment(line_id=line.id, comment_date=date(2026, 10, 18), text="Старое сообщение", author_username="old.master", author_name="Мастер")); db.commit()
    path = f"/lines/{line.id}/comments/2026-10-18"
    first = client.post(path, json={"text": "Первый ответ"})
    assert first.status_code == 200
    assert first.json()["comment"]["author_username"] == "regular.admin"
    actor = db.scalar(select(User).where(User.username == "regular.admin"))
    actor.role = "planner"; actor.display_name = "Планер"; db.commit()
    assert client.post(path, json={"text": "Второй ответ"}).status_code == 200
    messages = client.get("/lines/insights?start=2026-10-18&end=2026-10-18").json()["comments"]
    assert {row["text"] for row in messages} == {"Старое сообщение", "Первый ответ", "Второй ответ"}
    assert len({row["id"] for row in messages}) == 3
    assert next(row for row in messages if row["text"] == "Второй ответ")["author_name"] == "Планер"
    assert all(row["created_at"] for row in messages)
    assert db.scalar(select(func.count(LineCommentMessage.id))) == 2
    assert client.post(path, json={"text": "   "}).status_code == 422
    assert client.post(path, json={"text": "x" * 2001}).status_code == 422
    assert db.scalar(select(func.count(LineCommentMessage.id))) == 2
    assert db.scalar(select(LineComment)).text == "Старое сообщение"


def test_line_conversation_obeys_writer_permissions_and_date_scope(client, db):
    line, other = list(db.scalars(select(ProductionLine).limit(2)))
    actor = db.scalar(select(User).where(User.username == "regular.admin"))
    actor.role = "master"; actor.line_name = line.name; actor.workshop_code = line.workshop_code; db.commit()
    assert client.post(f"/lines/{line.id}/comments/2026-10-18", json={"text": "Сообщение мастера"}).status_code == 200
    assert client.post(f"/lines/{other.id}/comments/2026-10-18", json={"text": "Чужая линия"}).status_code == 403
    actor.role = "viewer"; db.commit()
    assert client.post(f"/lines/{line.id}/comments/2026-10-18", json={"text": "Просмотр"}).status_code == 403
    assert len(client.get("/lines/insights?start=2026-10-18&end=2026-10-18").json()["comments"]) == 1
    assert not client.get("/lines/insights?start=2026-10-19&end=2026-10-19").json()["comments"]


@pytest.mark.parametrize("role,visible", [("admin", True), ("planner", True), ("master", False), ("viewer", False)])
def test_only_planners_receive_unplaced_diagnostics(client, db, role, visible):
    plan, demand = make_plan(db)
    for row in db.scalars(select(LineCapability).where(LineCapability.product_id == demand.product_id)):
        db.delete(row)
    db.flush(); PlanService(db).calculate(plan, [demand]); db.commit()
    actor = db.scalar(select(User).where(User.username == "regular.admin"))
    actor.role = role; actor.line_name = "Сэндвичи"; db.commit()
    response = client.get("/plans/active/matrix?start=2026-09-08&days=1")
    assert response.status_code == 200
    assert bool(response.json()["unscheduled"]) is visible


def test_conversation_migration_preserves_legacy_comments_and_allows_multiple_messages():
    import importlib.util
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    migration_path = Path(__file__).parents[1] / "alembic/versions/20261018_0019_line_comment_messages.py"
    spec = importlib.util.spec_from_file_location("conversation_migration", migration_path)
    migration = importlib.util.module_from_spec(spec); spec.loader.exec_module(migration)
    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        connection.exec_driver_sql("CREATE TABLE production_lines (id INTEGER PRIMARY KEY)")
        connection.exec_driver_sql("INSERT INTO production_lines VALUES (1)")
        connection.exec_driver_sql("CREATE TABLE line_comments (id INTEGER PRIMARY KEY, text TEXT)")
        connection.exec_driver_sql("INSERT INTO line_comments VALUES (1, 'legacy')")
        with Operations.context(MigrationContext.configure(connection)):
            migration.upgrade()
        for content in ("first", "second"):
            connection.exec_driver_sql("INSERT INTO line_comment_messages(line_id,comment_date,text,author_username,author_name) VALUES (1,'2026-10-18',?,'test','Test')", (content,))
        assert connection.exec_driver_sql("SELECT text FROM line_comments").scalar() == "legacy"
        assert connection.exec_driver_sql("SELECT COUNT(*) FROM line_comment_messages").scalar() == 2
    engine.dispose()


def test_zam_dates_match_production_after_create_and_move(client, db):
    plan, demand = make_plan(db)
    line = plan.schedule_items[0].line
    response = client.post(f"/plans/{plan.id}/items", headers={"If-Match": str(plan.revision)}, json={"product_id": demand.product_id, "line_id": line.id, "production_date": "2026-09-08", "marking_date": "2026-09-20", "source_kind": "zam", "quantity": 10})
    assert response.status_code == 200, response.text
    item = next(row for row in plan.schedule_items if row.source_kind == "zam")
    assert item.marking_date == item.production_date
    PlanService(db).update_item(item, {"production_date": date(2026, 9, 9), "marking_date": date(2026, 10, 1)})
    result = schedule_item_dict(item)
    assert result["production_date"] == result["marking_date"] == result["shipping_date"] == date(2026, 9, 9)
    from app.services.notification_service import _shipping_date
    from app.services.export_service import ExcelExportService
    assert _shipping_date(result) == ExcelExportService._shipping_date(result) == date(2026, 9, 9)


def test_ohl_overload_is_displayed_and_included_in_plan_email(client, db):
    plan, demand = make_plan(db, quantity=250)
    ohl = [row for row in plan.schedule_items if row.schedule_kind == "production"]
    assert sum(row.quantity for row in ohl) == 250
    assert all(row.production_date == demand.requested_date for row in ohl)
    assert all(row.status.value != "conflict" for row in ohl)
    assert any(row.load_percent > 100 for row in ohl)
    response = client.get("/admin/mail-preview?start=2026-09-08&end=2026-09-08")
    assert response.status_code == 200
    assert response.json()["item_count"] == len(ohl)


def test_pc_manual_order_and_wash_duration_survive_recalculation(db):
    day = date(2026, 10, 22)
    line = db.scalar(select(ProductionLine).where(ProductionLine.workshop_code == "PC"))
    first = Product(sku="PC-OHL", name="Булочка", mono_group="Булочка")
    second = Product(sku="PC-OTHER", name="Слойка", mono_group="Слойка")
    plan = ProductionPlan(name="PC manual order", horizon_start=day, horizon_end=day)
    db.add_all([first, second, plan]); db.flush()
    for product, kind in ((first, "ohl"), (second, "generic")):
        db.add(LineCapability(line_id=line.id, product_id=product.id, units_per_hour=100))
        plan.schedule_items.append(ProductionScheduleItem(product_id=product.id, line_id=line.id, production_date=day, source_kind=kind, source="manual", schedule_kind="production", quantity=100, quantity_kg=100, required_hours=1))
    db.flush(); service = PlanService(db); service.refresh_sequence_and_cleanings(plan); db.commit()
    before = next(row for row in plan.schedule_items if row.product_id == first.id)
    other = next(row for row in plan.schedule_items if row.product_id == second.id)
    service.update_item(other, {"before_item_id": before.id})
    assert other.sequence < before.sequence
    washes = [row for row in plan.schedule_items if row.schedule_kind == "cleaning" and not row.excluded]
    assert len(washes) == 2
    wash = washes[0]
    service.update_item(wash, {"duration_minutes": 45, "sequence": 1})
    assert wash.sequence == 1 and wash.required_hours == Decimal("0.75")
    assert not wash.start_time and not wash.end_time
    service.calculate(plan, [])
    assert wash.sequence == 1 and wash.duration_hours == Decimal("0.75")
    assert other.sequence < before.sequence
    assert len([row for row in plan.schedule_items if row.schedule_kind == "cleaning" and not row.excluded]) == 2


def test_manual_pf_is_valid_source_and_unknown_source_is_rejected(client, db):
    plan, demand = make_plan(db)
    service = PlanService(db)
    values = {"product_id": demand.product_id, "line_id": plan.schedule_items[0].line_id, "production_date": plan.horizon_start, "quantity": 10, "source_kind": "pf"}
    service.create_manual_task(plan, values)
    assert any(row.source_kind == "pf" and row.quantity == 10 for row in plan.schedule_items)
    with pytest.raises(ValueError):
        service.create_manual_task(plan, {**values, "source_kind": "unknown"})


def test_complete_ohl_upgrade_is_once_and_preserves_unchanged_execution_metadata(db):
    plan, demand = make_plan(db, 50)
    original = next(row for row in plan.schedule_items if row.schedule_kind == "production")
    original.execution_status = "exported"; original.execution_note = "Existing export"; db.commit()
    service = PlanService(db)
    assert service.apply_complete_ohl_upgrade() is True
    current = next(row for row in plan.schedule_items if row.schedule_kind == "production")
    assert current.execution_status == "exported" and current.execution_note == "Existing export"
    revision = plan.revision
    assert service.apply_complete_ohl_upgrade() is False
    assert plan.revision == revision
