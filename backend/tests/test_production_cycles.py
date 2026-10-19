from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from test_audit_regressions import db, client
from test_line_import_scope import two_lines, snapshot
from app.api.routes import lines
from app.models.entities import ProductionLine, ProductionScheduleItem, User
from app.services.production_cycle_service import cycle_window, validate_cycle, production_cycle
from app.services.plan_service import PlanService, schedule_item_dict
from app.services.line_schedule_service import shift_hours
from app.services.notification_service import build_plan_xlsx
from app.services.export_service import ExcelExportService
from app.services.planning_engine import PlanningEngine, DemandInput, CapabilityInput, CapacityInput


CYCLE = {"enabled": True, "start_time": "18:00", "duration_hours": 24}


@pytest.mark.parametrize("day,start,hours,expected", [
    (date(2026, 10, 19), "18:00", 24, (datetime(2026, 10, 18, 18), datetime(2026, 10, 19, 18))),
    (date(2027, 1, 1), "18:30", 24, (datetime(2026, 12, 31, 18, 30), datetime(2027, 1, 1, 18, 30))),
    (date(2026, 10, 19), "08:00", 12, (datetime(2026, 10, 19, 8), datetime(2026, 10, 19, 20))),
    (date(2026, 10, 19), "23:45", 1.5, (datetime(2026, 10, 18, 23, 45), datetime(2026, 10, 19, 1, 15))),
])
def test_cycle_calendar_bounds(day, start, hours, expected):
    line = ProductionLine(planning_settings={"production_cycle": {**CYCLE, "start_time": start, "duration_hours": hours}})
    assert cycle_window(line, day) == expected
    line.planning_settings = {}
    assert cycle_window(line, day) is None


@pytest.mark.parametrize("bad", [
    {"enabled": "yes"}, {"start_time": "24:00"}, {"start_time": "18:70"},
    {"duration_hours": 0}, {"duration_hours": 25}, {"duration_hours": "NaN"},
    {"duration_hours": 1.001},
])
def test_invalid_cycle(bad):
    with pytest.raises(ValueError):
        validate_cycle({**CYCLE, **bad})


def test_same_dm_uses_shared_capacity_after_ohl():
    dm = date(2026, 10, 19)
    cap = CapabilityInput(1, 1, Decimal(100), marking_day_cycle=True)
    demands = [DemandInput(1, 1, "OHL", Decimal(150), date(2026, 10, 18), date(2026, 10, 18),
                           source_kind="ohl", marking_date=dm),
               DemandInput(2, 1, "ZAM", Decimal(100), dm, dm, source_kind="zam")]
    slots = [CapacityInput(1, dm, Decimal(2))]
    result = PlanningEngine().plan(demands, [cap], slots, dm, complete_zam=False)
    assert sum(row.quantity for row in result if row.source_kind == "ohl" and row.production_date == dm) == 150
    assert sum(row.quantity for row in result if row.source_kind == "zam" and row.production_date == dm) == 50
    assert sum(row.quantity for row in result if row.source_kind == "zam" and row.production_date is None) == 50
    full = PlanningEngine().plan(demands, [cap], slots, dm, complete_zam=True)
    assert sum(row.quantity for row in full if row.source_kind == "zam" and row.production_date == dm) == 100
    assert all(row.marking_date == row.production_date for row in full if row.source_kind == "zam")


def test_settings_save_rebuilds_cycle_and_preserves_other_lines(client, db, monkeypatch):
    monkeypatch.setattr(lines, "send_notification", lambda *args, **kwargs: None)
    plan, first, second, target, other = two_lines(db, monkeypatch)
    first.product.advance_status = "АЗ"
    first.marking_date = first.source_plan_date = date(2026, 9, 9)
    db.commit()
    PlanService(db).calculate(plan, PlanService(db)._latest_source_demands())
    old = next(row for row in plan.schedule_items if row.demand_item_id == first.id and row.schedule_kind == "production")
    old.execution_status = "exported"; old.execution_note = "keep me"; db.commit()
    protected = snapshot(plan, other)
    payload = {**client.get(f"/lines/{target}/schedule").json(), "production_cycle": CYCLE, "plan_version": plan.revision}
    payload["slots"] = []
    payload["schedule_code"] = "two_shift_daily"
    response = client.put(f"/lines/{target}/schedule", json=payload)
    assert response.status_code == 200, response.text
    assert snapshot(plan, other) == protected
    item = next(row for row in plan.schedule_items if row.demand_item_id == first.id and row.schedule_kind == "production")
    assert item.production_date == first.marking_date == date(2026, 9, 9)
    assert (item.execution_status, item.execution_note) == ("exported", "keep me")
    serialized = schedule_item_dict(item)
    assert serialized["production_window_start"] == datetime(2026, 9, 8, 18)
    assert serialized["production_window_end"] == datetime(2026, 9, 9, 18)
    got = client.get(f"/lines/{target}/schedule?start=2026-09-09").json()
    assert got["production_cycle"] == CYCLE
    assert got["slots"][0]["production_window_start"] == "2026-09-08T18:00:00"
    # The export preserves a real Excel datetime, rather than a date-only string.
    for content in (build_plan_xlsx([serialized], "production"), ExcelExportService().build([serialized])):
        book = load_workbook(BytesIO(content))
        sheet = book["План выгрузки"] if "План выгрузки" in book.sheetnames else book.active
        headings = [cell.value for cell in sheet[1]]
        assert sheet.cell(2, headings.index("Начало производственных суток") + 1).value == datetime(2026, 9, 8, 18)
    # A stale editor cannot overwrite the newer plan.
    assert client.put(f"/lines/{target}/schedule", json=payload).status_code == 409
    payload["plan_version"] = plan.revision
    payload["production_cycle"] = {**CYCLE, "enabled": False}
    assert client.put(f"/lines/{target}/schedule", json=payload).status_code == 200
    assert snapshot(plan, other) == protected
    item = next(row for row in plan.schedule_items if row.demand_item_id == first.id and row.schedule_kind == "production")
    assert item.production_date == date(2026, 9, 8)  # legacy AZ convention restored
    assert schedule_item_dict(item)["production_window_start"] is None
    db.scalar(select(User)).role = "viewer"; db.commit()
    assert client.put(f"/lines/{target}/schedule", json=payload).status_code == 403


def test_window_caps_generated_hours_and_rejects_oversized_manual_slots(client, db, monkeypatch):
    monkeypatch.setattr(lines, "send_notification", lambda *args, **kwargs: None)
    line = db.scalar(select(ProductionLine).where(ProductionLine.name == "Сэндвичи"))
    line.schedule_code = "two_shift_daily"
    line.planning_settings = {"production_cycle": {**CYCLE, "duration_hours": 8}}
    db.commit()
    assert shift_hours(line, date(2026, 10, 19)) == (4, 4)
    payload = client.get(f"/lines/{line.id}/schedule").json()
    payload["slots"] = [{"capacity_date": "2026-10-19", "day_hours": 11, "night_hours": 11}]
    assert client.put(f"/lines/{line.id}/schedule", json=payload).status_code == 422
    payload["slots"] = []; payload["production_cycle"] = {**CYCLE, "duration_hours": "wrong"}
    assert client.put(f"/lines/{line.id}/schedule", json=payload).status_code == 422


def test_upgrade_is_once_and_does_not_change_other_lines(db):
    before = {row.id: dict(row.planning_settings or {}) for row in db.scalars(select(ProductionLine))}
    service = PlanService(db)
    assert service.apply_production_cycle_upgrade()
    line = db.scalar(select(ProductionLine).where(ProductionLine.name == "Сэндвичи"))
    assert production_cycle(line) == CYCLE
    for row in db.scalars(select(ProductionLine)):
        if row.id != line.id:
            assert row.planning_settings == before[row.id]
    line.planning_settings = {"production_cycle": {**CYCLE, "enabled": False}}; db.commit()
    assert not service.apply_production_cycle_upgrade()
    assert not production_cycle(line)["enabled"]


def test_shared_sku_on_another_line_is_reserved_not_replaced(db, monkeypatch):
    plan, first, _, target, other = two_lines(db, monkeypatch)
    split = ProductionScheduleItem(plan_id=plan.id, product_id=first.product_id, demand_item_id=first.id,
        line_id=other, production_date=first.requested_date, quantity=20, quantity_kg=20,
        required_hours=Decimal("0.2"), source_kind="ohl", source="auto", shift="day")
    plan.schedule_items.append(split); db.commit()
    before = snapshot(plan, other)
    PlanService(db).calculate(plan, [first], "line_schedule_changed", line_id=target, preserve_other_lines=True)
    assert snapshot(plan, other) == before
    assert sum(row.quantity for row in plan.schedule_items if row.demand_item_id == first.id and row.schedule_kind == "production") == 100
