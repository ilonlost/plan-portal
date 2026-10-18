from datetime import date
from io import BytesIO
from openpyxl import load_workbook

from app.core.auth_service import _role
from app.core.config import settings
from app.services.notification_service import build_plan_email_html, build_plan_xlsx
from app.services.settings_service import default_mail_configuration


def test_plan_email_contains_units_and_line_grouping() -> None:
    html = build_plan_email_html(
        default_mail_configuration(),
        "План 33 недели",
        date(2026, 8, 27),
        date(2026, 8, 27),
        [{
            "production_date": "2026-08-27", "line_name": "Бургеры", "sku": "1010033732",
            "product_name": "Чизбургер", "shift": "day", "quantity_units": 120,
            "quantity_kg": 72.6, "required_hours": 1.25, "source_kind": "ohl",
            "marking_date": "2026-08-28", "shipping_date": "2026-08-27",
        }],
    )
    assert "Бургеры" in html
    assert "120 шт." in html
    assert "72.60 кг" in html
    assert "1010033732" in html
    assert "Дата отгрузки" in html and "27.08.2026" in html


def test_plan_email_workbook_contains_shipping_date_column():
    content = build_plan_xlsx([{
        "production_date": "2026-08-27", "marking_date": "2026-08-28", "shipping_date": "2026-08-27",
        "line_name": "Бургеры", "sku": "1010033732", "product_name": "Чизбургер", "quantity_kg": 72.6,
    }], "production")
    sheet = load_workbook(BytesIO(content), read_only=True).active
    assert list(next(sheet.iter_rows(values_only=True)))[:3] == ["Дата производства", "Дата маркировки", "Дата отгрузки"]
    assert list(next(sheet.iter_rows(min_row=2, values_only=True)))[:3] == ["27.08.2026", "28.08.2026", "27.08.2026"]


def test_shipping_date_fallback_subtracts_day_across_month_boundary():
    content = build_plan_xlsx([{
        "production_date": "2026-08-31", "marking_date": "2026-09-01",
        "line_name": "Бургеры", "sku": "1010033732", "product_name": "Чизбургер",
    }], "production")
    sheet = load_workbook(BytesIO(content), read_only=True).active
    assert list(next(sheet.iter_rows(min_row=2, values_only=True)))[:3] == ["31.08.2026", "01.09.2026", "31.08.2026"]


def test_ldap_non_admin_roles_are_managed_in_web_admin(monkeypatch) -> None:
    monkeypatch.setattr(settings, "portal_admin_logins", "root.admin")
    monkeypatch.setattr(settings, "ldap_admin_group_dn", "CN=PLAN-Admins")
    monkeypatch.setattr(settings, "ldap_planner_group_dn", "CN=PLAN-Planners")
    monkeypatch.setattr(settings, "ldap_master_group_dn", "CN=PLAN-Masters")
    assert _role("root.admin", []) == "admin"
    assert _role("planner.user", ["CN=PLAN-Planners"]) == "viewer"
    assert _role("master.user", ["CN=PLAN-Masters"]) == "viewer"
