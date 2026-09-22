from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from app.core.config import settings
from app.services import tail_buffer_service as service
from app.services import production_fact_export as export
from tests.test_audit_regressions import client, db


def row(index=0):
    return dict(id=str(index), record_id=str(index), source_month="202609", sscc=f"001234567890{index:08d}",
                source_center="5410", target_buffer="5498", workshop_code="KC", line_name="Жареные блюда",
                moved_at=datetime(2026, 9, 16, 13, 45, tzinfo=service.MOSCOW),
                first_moved_at=datetime(2026, 9, 16, 12, tzinfo=service.MOSCOW),
                sku="001010017483", product_name="=SUM(A1:A3)", created_date=date(2026, 9, 15),
                buffer_name="Буфер", card_buffer=5498, warning="", weight_kg=Decimal("23.75"), shelf_life_days=Decimal("14"))


@pytest.fixture(autouse=True)
def no_live_cycles(monkeypatch):
    monkeypatch.setattr(export, "iter_cycle_rows", lambda *args: iter([]))


def test_export_all_rows_identifiers_dates_text_and_summary(client, monkeypatch):
    calls = []
    def records(start, end, center, search):
        calls.append((start, end, center, search))
        return (row(index) for index in range(5101))
    monkeypatch.setattr(export, "iter_export_rows", records)
    response = client.get("/production-fact/export.xlsx?start=2026-09-16&end=2026-09-16&center=5410&search=сырники")
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Сводка", "Проводки", "Почасовой выпуск", "Циклы"]
    sheet = book["Проводки"]
    assert sheet.max_row == 5102 and sheet.auto_filter.ref == "A1:Q5102"
    assert sheet.freeze_panes == "A2"
    assert sheet["J2"].value == "00123456789000000000" and sheet["J2"].data_type == "s"
    assert sheet["E2"].value == "001010017483" and sheet["F2"].data_type == "s"
    assert sheet["A2"].value == datetime(2026, 9, 16, 13, 45)
    assert sheet["P2"].value == 23.75 and sheet["P2"].data_type == "n"
    assert sheet["Q1"].value == "Срок годности, суток"
    assert sheet["Q2"].value == 14 and sheet["Q2"].data_type == "n"
    assert book["Почасовой выпуск"]["E15"].value == 5101 * 23.75
    assert sheet["G2"].value == datetime(2026, 9, 15)
    assert book["Сводка"]["B7"].value == 5101
    assert book["Сводка"]["D14"].value == 5101
    assert calls == [(date(2026, 9, 16), date(2026, 9, 16), "5410", "сырники")]


def test_export_sheet_rollover_and_empty(monkeypatch):
    monkeypatch.setattr(export, "MAX_DATA_ROWS", 2)
    monkeypatch.setattr(export, "iter_export_rows", lambda *args: (row(i) for i in range(5)))
    path = export.export_workbook(date(2026, 9, 16), date(2026, 9, 16))
    try:
        book = load_workbook(path)
        assert [book[name].max_row for name in book.sheetnames if name.startswith("Проводки")] == [3, 3, 2]
        assert book["Сводка"]["B7"].value == 5
        book.close()
    finally:
        Path(path).unlink()
    monkeypatch.setattr(export, "iter_export_rows", lambda *args: iter([]))
    path = export.export_workbook(date(2026, 9, 16), date(2026, 9, 16))
    try:
        book = load_workbook(path)
        assert book["Проводки"].max_row == 1 and book["Сводка"]["B7"].value == 0
        book.close()
    finally:
        Path(path).unlink()


def test_export_refuses_partial_source_and_validates_filters(client, monkeypatch):
    def broken(*args):
        yield row()
        raise service.SourceError("partial", ["Месяц недоступен"])
    monkeypatch.setattr(export, "iter_export_rows", broken)
    response = client.get("/production-fact/export.xlsx")
    assert response.status_code == 503 and "Месяц недоступен" in response.json()["detail"]
    assert client.get("/production-fact/tail-buffers?center=evil").status_code == 422
    assert client.get("/production-fact/tail-buffers?offset=-1").status_code == 422
    assert client.get("/production-fact/tail-buffers?limit=5000").status_code == 422


def test_summary_includes_all_rows_and_empty_centers(monkeypatch):
    engine = MagicMock()
    connection = engine.connect.return_value.__enter__.return_value
    present = MagicMock(); present.scalar.return_value = 1
    records = MagicMock()
    records.mappings.return_value.all.return_value = [
        dict(source_center=None, bucket=None, all_centers=1, is_total=1, postings=12000, sscc_count=11000, weight_kg=Decimal("15000.25"), missing_weight_count=3),
        dict(source_center=5410, bucket=None, all_centers=0, is_total=1, postings=12000, sscc_count=11000, weight_kg=Decimal("15000.25"), missing_weight_count=3, first_at=datetime(2026, 9, 16), last_at=datetime(2026, 9, 16, 1)),
        dict(source_center=5410, bucket=datetime(2026, 9, 16), all_centers=0, is_total=0, postings=12000, sscc_count=11000, weight_kg=Decimal("15000.25"), missing_weight_count=3),
    ]
    connection.execute.side_effect = [present, records]
    monkeypatch.setattr(settings, "production_fact_database_url", "mssql+pymssql://u:p@s/db")
    monkeypatch.setattr(service, "create_engine", MagicMock(return_value=engine))
    report = service.load_summary(date(2026, 9, 16), date(2026, 9, 16))
    assert report["status"] == "connected" and report["total_count"] == 12000
    assert report["weight_kg"] == Decimal("15000.25") and report["missing_weight_count"] == 3
    assert report["centers"][1]["buckets"][0]["weight_kg"] == Decimal("15000.25")
    assert report["active_centers"] == 1 and report["sscc_count"] == 11000
    assert len(report["centers"]) == 12 and report["bucket_minutes"] == 60
    assert sum(c["postings"] for c in report["centers"]) == report["total_count"]


def test_full_export_query_has_no_pagination_and_literal_search():
    sql = service.build_query("dbo", ["cp_DWH_LA0052_202609"], export=True, searching=True)
    assert "OFFSET" not in sql and "TOP" not in sql and ":limit" not in sql
    assert "LIKE :search" in sql and "TRY_CONVERT(int, :center)" in sql
    assert service.filter_params("5410", "50%_[") == {"center": "5410", "search": "%50~%~_~[%"}


def test_cycles_exact_five_minutes_midnight_lines_negative_and_unknown_weight():
    at = datetime(2026, 9, 16, 23, 58)
    records = [
        dict(source_center=5410, moved_at=at, weight_kg=Decimal("10.25")),
        dict(source_center=5410, moved_at=at + timedelta(minutes=4, seconds=59), weight_kg=Decimal("-0.25")),
        dict(source_center=5410, moved_at=at + timedelta(minutes=9, seconds=59), weight_kg=Decimal("5")),
        dict(source_center=5410, moved_at=at + timedelta(minutes=9, seconds=59), weight_kg=None),
        dict(source_center=5810, moved_at=at, weight_kg=None),
    ]
    cycles = list(service.group_cycles(records))
    assert len(cycles) == 3
    assert cycles[0]["postings"] == 2 and cycles[0]["weight_kg"] == 10
    assert cycles[0]["end_at"].day == 17 and cycles[0]["gap_before_minutes"] is None
    assert cycles[1]["gap_before_minutes"] == 5 and cycles[1]["postings"] == 2
    assert cycles[1]["missing_weight_count"] == 1
    assert cycles[2]["weight_kg"] is None and cycles[2]["gap_before_minutes"] is None
    assert list(service.group_cycles([])) == []


def test_export_hours_preserve_weight_and_cycles_ignore_article_search(monkeypatch):
    at = datetime(2026, 9, 16, 13, 59, tzinfo=service.MOSCOW)
    records = [{**row(), "moved_at": at, "weight_kg": Decimal("10.25")},
               {**row(1), "moved_at": at + timedelta(minutes=1), "weight_kg": Decimal("20.5")},
               {**row(2), "moved_at": at + timedelta(minutes=6), "weight_kg": None}]
    monkeypatch.setattr(export, "iter_export_rows", lambda *args: iter(records))
    calls = []
    def cycles(start, end, center):
        calls.append((start, end, center))
        return service.group_cycles(records)
    monkeypatch.setattr(export, "iter_cycle_rows", cycles)
    path = export.export_workbook(date(2026, 9, 16), date(2026, 9, 16), "5410", "Артикул")
    try:
        book = load_workbook(path)
        hourly = book["Почасовой выпуск"]
        assert hourly["E15"].value == 10.25 and hourly["E16"].value == 20.5
        assert hourly["H16"].value == 1
        assert hourly.max_row == 25 and hourly["E2"].value == 0
        assert hourly["C15"].value == datetime(2026, 9, 16, 13)
        sheet = book["Циклы"]
        assert sheet.max_row == 3 and sheet["F2"].value == 30.75
        assert sheet["G2"].value == 1845  # 30.75 kg / one-minute posting interval
        assert sheet["H3"].value == 5 and sheet["G3"].value is None
        assert calls == [(date(2026, 9, 16), date(2026, 9, 16), "5410")]
        book.close()
    finally:
        Path(path).unlink()
