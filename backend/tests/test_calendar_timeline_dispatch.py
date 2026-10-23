from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from test_audit_regressions import db, client, make_plan
from app.api.routes import lines, plans
from app.models.entities import ProductionLine, ProductionScheduleItem, LineCapability, User
from app.services.plan_service import PlanService, schedule_item_dict
from app.services.task_timeline_service import bind_task_timeline
from app.services.export_service import ExcelExportService
from app.services.notification_service import build_plan_xlsx, build_plan_email_html
from app.core.config import settings
from app.core.security import UserContext, create_session_token


def test_sun_moon_sun_calendar_dates_without_moving_the_plan_day(db):
    plan, demand = make_plan(db)
    first = next(r for r in plan.schedule_items if r.schedule_kind == "production")
    first.line.planning_settings = {"production_cycle":{"enabled":True,"start_time":"18:00","duration_hours":24}}
    first.production_date = date(2026,10,22); first.marking_date = date(2026,10,23); first.required_hours = 2; first.sequence = 1
    for sequence,hours in [(2,12),(3,10)]:
        plan.schedule_items.append(ProductionScheduleItem(line=first.line, product=first.product, production_date=first.production_date,
            marking_date=first.marking_date, sequence=sequence, required_hours=hours, quantity=10, quantity_kg=10, source_kind="ohl", schedule_kind="production", excluded=False))
    db.flush(); bind_task_timeline(plan.schedule_items)
    rows = [schedule_item_dict(r) for r in sorted(plan.schedule_items,key=lambda r:r.sequence)]
    assert [(r['task_start_at'],r['task_end_at']) for r in rows] == [
        (datetime(2026,10,21,18),datetime(2026,10,21,20)),
        (datetime(2026,10,21,20),datetime(2026,10,22,8)),
        (datetime(2026,10,22,8),datetime(2026,10,22,18))]
    assert [r['shift_segments'][0]['shift'] for r in rows] == ['day','night','day']
    assert [r['calendar_production_start'] for r in rows] == [date(2026,10,21),date(2026,10,21),date(2026,10,22)]
    assert rows[1]['calendar_production_end'] == date(2026,10,22)
    assert all(r['production_date'] == date(2026,10,22) and r['marking_date'] == date(2026,10,23) and r['shipping_date'] == date(2026,10,22) and r['shipping_time'] == time(19) for r in rows)
    # A long single card exposes every day/night segment rather than a false fixed shift.
    first.required_hours = 24; bind_task_timeline([first])
    assert [r['shift'] for r in schedule_item_dict(first)['shift_segments']] == ['day','night','day']


def test_dispatch_override_survives_move_rebuild_and_all_exports(client, db, monkeypatch):
    monkeypatch.setattr(plans,"send_notification",lambda *a,**k:None)
    plan,demand = make_plan(db)
    item = next(r for r in plan.schedule_items if r.schedule_kind=='production')
    result = client.patch(f'/plans/{plan.id}/items/{item.id}',headers={'If-Match':str(plan.revision)},json={'shipping_date':'2026-10-25','shipping_time':'19:30'})
    assert result.status_code == 200,result.text
    PlanService(db).update_item(item,{'production_date':date(2026,10,24)})
    PlanService(db).calculate(plan,[demand])
    data = schedule_item_dict(item)
    assert data['shipping_date'] == date(2026,10,25) and data['shipping_time'] == time(19,30) and data['shipping_date_overridden']
    data['source_kind']='zam' # Dispatch overrides must also win in ZAM-specific export fallbacks.
    for content in [ExcelExportService().build([data]),build_plan_xlsx([data],'production')]:
        book=load_workbook(BytesIO(content)); sheet=book['План выгрузки'] if 'План выгрузки' in book.sheetnames else book.active
        headings=[c.value for c in sheet[1]]
        assert str(sheet.cell(2,headings.index('Дата отгрузки')+1).value).startswith(('2026-10-25','25.10.2026'))
        assert sheet.cell(2,headings.index('Отгрузка до')+1).value == time(19,30)
    html=build_plan_email_html({},plan.name,date(2026,10,24),date(2026,10,24),[data],'production')
    assert '25.10.2026 до 19:30' in html
    result = client.patch(f'/plans/{plan.id}/items/{item.id}',headers={'If-Match':str(plan.revision)},json={'shipping_date':None,'shipping_time':None})
    assert result.status_code==200
    assert not schedule_item_dict(item)['shipping_date_overridden']


def test_lasagna_reorder_can_cross_source_and_shift_and_survives_recalculation(client,db,monkeypatch):
    monkeypatch.setattr(plans,'send_notification',lambda *a,**k:None)
    plan,demand=make_plan(db)
    line=db.scalar(select(ProductionLine).where(ProductionLine.name=='Лазанья'))
    db.add(LineCapability(line_id=line.id,product_id=demand.product_id,units_per_hour=100));db.commit()
    for source,shift in [('ohl','day'),('zam','night'),('pf','day')]:
        result=client.post(f'/plans/{plan.id}/items',headers={'If-Match':str(plan.revision)},json={'product_id':demand.product_id,'line_id':line.id,'production_date':'2026-10-22','marking_date':'2026-10-23','quantity':10,'source_kind':source,'shift':shift})
        assert result.status_code==200,result.text
    rows=sorted([r for r in plan.schedule_items if r.line_id==line.id and r.schedule_kind=='production'],key=lambda r:r.sequence)
    first,last=rows[0],rows[-1]
    result=client.patch(f'/plans/{plan.id}/items/{last.id}',headers={'If-Match':str(plan.revision)},json={'before_item_id':first.id,'production_date':'2026-10-22'})
    assert result.status_code==200,result.text
    assert last.sequence < first.sequence
    PlanService(db).calculate(plan,[demand])
    assert last.sequence < first.sequence
    matrix=client.get(f'/plans/active/matrix?start=2026-10-19&days=7&line_id={line.id}').json()
    ids=[i['id'] for w in matrix['workshops'] for l in w['lines'] for c in l['cells'] if c['date']=='2026-10-22' for i in c['items'] if i['schedule_kind']=='production']
    assert ids.index(last.id) < ids.index(first.id)


def test_corporate_viewer_can_start_and_reply_without_plan_write_access(client,db,monkeypatch):
    monkeypatch.setattr(lines,'send_notification',lambda *a,**k:None)
    plan,_=make_plan(db)
    actor=db.scalar(select(User).where(User.username=='regular.admin'));actor.role='viewer';db.commit()
    line=db.scalar(select(ProductionLine))
    assert client.get('/session/me').json()['can_comment']
    for message in ['Первое сообщение','Ответ']:
        assert client.post(f'/lines/{line.id}/comments/2026-10-22',json={'text':message}).status_code==200
    assert len(client.get('/lines/insights?start=2026-10-22&end=2026-10-22').json()['comments'])==2
    assert client.post(f'/plans/{plan.id}/events',headers={'If-Match':str(plan.revision)},json={'line_id':line.id,'production_date':'2026-10-22','schedule_kind':'cleaning','duration_hours':1}).status_code==403


def test_service_viewer_cannot_comment_even_with_viewer_permission(client,db):
    actor=db.scalar(select(User));actor.role='viewer';actor.auth_provider='service';db.commit()
    client.cookies.set(settings.session_cookie_name,create_session_token(UserContext(actor.username,actor.display_name,'viewer',auth_method='service')))
    assert not client.get('/session/me').json()['can_comment']
    line=db.scalar(select(ProductionLine))
    assert client.post(f'/lines/{line.id}/comments/2026-10-22',json={'text':'Не должно отправиться'}).status_code==403
